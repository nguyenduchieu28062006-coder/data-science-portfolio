"""Isolated, stateless sklearn benchmark. CSV and models exist only in memory.

Vercel entrypoint: handler. Local equivalent: python -B api/ml-compare.py --serve.
The HTTP handler runs training in a killable worker with a 55 second deadline.
"""
import csv
import base64
import zipfile
from xml.etree import ElementTree
import io
import json
import logging
import math
import os
import pickle
import re
import signal
import subprocess
import sys
import time
import warnings
from datetime import datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer, SimpleHTTPRequestHandler
from pathlib import Path

logger = logging.getLogger(__name__)
if not logger.handlers:
    _log_handler = logging.StreamHandler(sys.stderr)
    _log_handler.setFormatter(logging.Formatter("%(levelname)s %(name)s %(message)s"))
    logger.addHandler(_log_handler)
logger.setLevel(logging.INFO)
logger.propagate = False

for _key in ("OMP_NUM_THREADS", "OPENBLAS_NUM_THREADS", "MKL_NUM_THREADS"):
    os.environ[_key] = "1"
import numpy as np
from sklearn.base import clone
from sklearn.compose import ColumnTransformer
from sklearn.impute import SimpleImputer
from sklearn.preprocessing import OneHotEncoder, StandardScaler
from sklearn.pipeline import Pipeline
from sklearn.model_selection import train_test_split, StratifiedKFold, KFold, TimeSeriesSplit
from sklearn.linear_model import LogisticRegression, LinearRegression, Ridge, Lasso
from sklearn.neighbors import KNeighborsClassifier, KNeighborsRegressor
from sklearn.svm import SVC, SVR
from sklearn.tree import DecisionTreeClassifier, DecisionTreeRegressor
from sklearn.ensemble import RandomForestClassifier, RandomForestRegressor, GradientBoostingClassifier, GradientBoostingRegressor
from sklearn.naive_bayes import GaussianNB
from sklearn import metrics, datasets
from threadpoolctl import threadpool_limits

MAX_CSV_BYTES = 2 * 1024 * 1024
MAX_REQUEST_BYTES = 4 * 1024 * 1024
MAX_ROWS, MAX_COLUMNS, MAX_CLASSES = 20000, 100, 30
MAX_EXPANDED, MAX_CELLS = 500, 3000000
MISSING = {"", "na", "n/a", "null", "none", "nan", "-", "--"}
ID_NAME = re.compile(r"(^|_)(id|uuid|index|row_num)($|_)", re.I)
NUM = re.compile(r"^[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?$")
DATE = re.compile(r"^\d{4}-\d{2}-\d{2}(?:[T ].*)?$")
DATE_LIKE = re.compile(r"^(?:\d{4}[-/]\d{1,2}[-/]\d{1,2}|\d{1,2}[-/]\d{1,2}[-/]\d{4})(?:[T ].*)?$")
TIME_LIKE = re.compile(r"^(?:[01]?\d|2[0-3]):[0-5]\d(?::[0-5]\d)?$")


class LabError(Exception):
    def __init__(self, message, status=422):
        super().__init__(message)
        self.status = status


def number(value, numeric_format=None):
    if value is None:
        return None
    token = str(value).strip()
    if re.fullmatch(r"[+-]?\d{1,3}(?:,\d{3})+\.\d+(?:[eE][+-]?\d+)?", token):
        token = token.replace(",", "")
    elif re.fullmatch(r"[+-]?\d{1,3}(?:\.\d{3})+,\d+(?:[eE][+-]?\d+)?", token):
        token = token.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"[+-]?\d+,\d+(?:[eE][+-]?\d+)?", token):
        # A lone 1,234 may mean either a decimal or a thousands separator.
        if ambiguous_number(token) and numeric_format != "decimal_comma":
            return None
        token = token.replace(",", ".")
    if not NUM.fullmatch(token):
        return None
    result = float(token)
    return result if math.isfinite(result) and abs(result) <= 1e100 else None


def ambiguous_number(value):
    return bool(value is not None and re.fullmatch(r"[+-]?\d{1,3},\d{3}", str(value)))


def column_numeric_format(values):
    # Resolve lone comma triples only with consistent unambiguous evidence in
    # this column. Feature callers supply TRAIN values, never held-out values.
    tokens = [str(v).strip() for v in values if v is not None]
    comma = [v for v in tokens if re.fullmatch(r"[+-]?\d+,\d+(?:[eE][+-]?\d+)?", v)]
    evidence = [v for v in comma if not ambiguous_number(v)]
    conflicting = any("." in v and number(v) is not None for v in tokens)
    # Plain integers agree with either decimal convention and are not contrary
    # evidence. Invalid/missing tokens do not dictate the locale of valid data.
    return "decimal_comma" if len(evidence) >= 2 and not conflicting else None


def csv_dialect(sample):
    try:
        detected = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        if len(next(csv.reader(io.StringIO(sample), detected))) >= 2:
            return detected
    except (csv.Error, StopIteration):
        pass
    # A short/ragged table can defeat Sniffer. Compare actual parsed records,
    # respecting quotes, rather than counting separators inside cell values.
    candidates = []
    for delimiter in (",", ";", "\t", "|"):
        try:
            records = [row for row in csv.reader(io.StringIO(sample), delimiter=delimiter) if row and any(v.strip() for v in row)]
            width = len(records[0]) if records else 0
            if 2 <= width <= MAX_COLUMNS:
                aligned = sum(len(row) == width for row in records[1:])
                oversized = sum(len(row) > width for row in records[1:])
                candidates.append((aligned - oversized * 2, width, delimiter))
        except csv.Error:
            continue
    return max(candidates)[2] if candidates else ","


def date(value):
    if value is None or not DATE_LIKE.match(value):
        return None
    try:
        if DATE.match(value):
            return datetime.fromisoformat(value.replace("Z", "+00:00"))
        parts = re.split(r"[-/]", value)
        if len(parts) != 3:
            return None
        a, b, c = map(int, parts)
        if len(parts[0]) == 4:
            return datetime(a, b, c)
        if a > 12:
            return datetime(c, b, a)
        if b > 12:
            return datetime(c, a, b)
        # Ambiguous day/month must be converted to ISO by the user.
        return None
    except ValueError:
        return None


def date_key(value):
    # UTC for explicit offsets; naive ISO values are interpreted as UTC consistently.
    from datetime import timezone
    d = date(value)
    return d.replace(tzinfo=timezone.utc).timestamp() if d and d.tzinfo is None else d.timestamp() if d else None


def parse_csv(text, check_size=True, delimiter=None):
    if not isinstance(text, str):
        raise LabError("CSV phải là văn bản UTF-8.", 400)
    if check_size and len(text.encode("utf-8")) > MAX_CSV_BYTES:
        raise LabError("CSV vượt giới hạn 2 MiB.", 413)
    if "\x00" in text:
        raise LabError("CSV chứa ký tự không hợp lệ.", 400)
    text = text.lstrip("\ufeff")
    try:
        sample = text[:8192]
        dialect = delimiter or csv_dialect(sample)
        reader = csv.reader(io.StringIO(text), delimiter=dialect, strict=True) if isinstance(dialect, str) else csv.reader(io.StringIO(text), dialect, strict=True)
        leading_blank = 0
        raw_headers = next(reader)
        while not raw_headers or all(not v.strip() for v in raw_headers):
            leading_blank += 1
            raw_headers = next(reader)
        headers, used = [], set()
        for i, raw in enumerate(raw_headers):
            base = raw.strip() or f"unnamed_{i + 1}"
            name, suffix = base, 2
            while name in used:
                name, suffix = f"{base}_{suffix}", suffix + 1
            if len(name) > 120:
                raise LabError("Tên cột quá dài (tối đa 120 ký tự).", 400)
            used.add(name)
            headers.append(name)
        if not 2 <= len(headers) <= MAX_COLUMNS:
            raise LabError("CSV cần từ 2 đến 100 cột.", 400)
        rows, blank, duplicates, seen = [], leading_blank, 0, set()
        input_rows, short_rows = 0, 0
        for raw in reader:
            if not raw or all(not v.strip() for v in raw):
                blank += 1
                continue
            input_rows += 1
            if input_rows > MAX_ROWS:
                raise LabError("CSV vượt giới hạn 20.000 dòng.", 413)
            if len(raw) > len(headers):
                raise LabError("Dòng CSV có nhiều ô hơn số cột; kiểm tra dấu phân cách.", 400)
            short_rows += len(raw) < len(headers)
            raw += [""] * (len(headers) - len(raw))
            row = tuple(None if v.strip().lower() in MISSING else v.strip() for v in raw)
            if any(v and len(v) > 4096 for v in row):
                raise LabError("Một ô CSV vượt 4.096 ký tự.", 400)
            if row in seen:
                duplicates += 1
            else:
                seen.add(row)
                rows.append(row)
        if not rows:
            raise LabError("CSV không có dòng dữ liệu.", 400)
        comma_count = sum(v is not None and "," in v and number(v) is not None for row in rows for v in row)
        ambiguous_count = sum(ambiguous_number(v) for row in rows for v in row)
        notes = []
        separator = dialect if isinstance(dialect, str) else dialect.delimiter
        if comma_count:
            notes.append(f"Tệp dùng dấu {separator!r}; đã nhận diện {comma_count} ô số có decimal comma / dấu phân nhóm hàng nghìn rõ ràng.")
        if ambiguous_count:
            notes.append(f"{ambiguous_count} ô số có dấu phẩy mơ hồ (ví dụ 1,234); cần bằng chứng decimal comma nhất quán trong cùng cột, nếu không giữ categorical để tránh parse sai.")
        return headers, rows, {"input_rows": input_rows, "blank_rows": blank, "duplicates_removed": duplicates,
                              "short_rows_padded": short_rows, "headers_normalized": headers != raw_headers,
                              "delimiter": separator, "localized_numeric_cells": comma_count, "ambiguous_numeric_cells": ambiguous_count,
                              "warnings": notes}
    except (csv.Error, StopIteration, UnicodeError):
        raise LabError("Không đọc được CSV UTF-8 hợp lệ.", 400) from None


def profile(headers, rows, audit):
    columns = []
    for i, name in enumerate(headers):
        values = [r[i] for r in rows if r[i] is not None]
        n, unique = len(values), len(set(values))
        numeric_format = column_numeric_format(values)
        numeric_count = sum(number(v, numeric_format) is not None for v in values)
        ambiguous = sum(ambiguous_number(v) for v in values) if not numeric_format else 0
        kind = "numeric" if n and not ambiguous and numeric_count / n >= .9 else "categorical"
        potential_date = bool(n and sum(bool(DATE_LIKE.match(v) or TIME_LIKE.match(v)) for v in values) / n >= .9)
        ratio = unique / n if n else 0
        uuid = bool(n and sum(bool(re.fullmatch(r"[0-9a-fA-F]{8}(?:-[0-9a-fA-F]{4}){3}-[0-9a-fA-F]{12}", v)) for v in values) / n > .9)
        nums = [number(v, numeric_format) for v in values]
        monotone = bool(kind == "numeric" and all(v is not None and v.is_integer() for v in nums) and n >= 30 and
                        (all(a < b for a, b in zip(nums, nums[1:])) or all(a > b for a, b in zip(nums, nums[1:]))))
        id_like = bool(ID_NAME.search(name) or re.search(r"email|product_code|case_id|order_id", name, re.I) or uuid or ratio >= .98 and (kind == "categorical" and not potential_date or monotone))
        suggested = "classification" if kind == "categorical" or unique <= 2 or unique <= 20 and ratio <= .15 else "regression"
        columns.append({"name": name, "kind": kind, "unique": unique, "unique_ratio": ratio,
                        "missing": len(rows) - n, "malformed_numeric": n - numeric_count if kind == "numeric" else 0,
                        "ambiguous_numeric": ambiguous,
                        "date": potential_date, "id_like": id_like, "high_cardinality": kind == "categorical" and (unique > 50 or unique > 20 and ratio > .5) and not potential_date,
                        "suggested_task": suggested, "target_candidate": n >= 2 and unique > 1 and not id_like and not potential_date})
    return {"rows": len(rows), "columns": columns, "audit": audit,
            "missing": sum(c["missing"] for c in columns), "numeric": sum(c["kind"] == "numeric" for c in columns),
            "categorical": sum(c["kind"] == "categorical" for c in columns), "preview": [list(r) for r in rows[:12]]}


def load_data(payload):
    if payload.get("demo"):
        names = {"iris": datasets.load_iris, "wine": datasets.load_wine, "breast_cancer": datasets.load_breast_cancer,
                 "digits": datasets.load_digits, "diabetes": datasets.load_diabetes}
        if payload["demo"] not in names:
            raise LabError("Dataset mẫu không hợp lệ.", 400)
        bunch = names[payload["demo"]]()
        headers = [str(n) for n in bunch.feature_names] + ["target"]
        rows = [tuple(str(float(v)) for v in x) + (str(float(y)),) for x, y in zip(bunch.data, bunch.target)]
        return headers, rows, {"input_rows": len(rows), "blank_rows": 0, "duplicates_removed": 0, "short_rows_padded": 0, "headers_normalized": False}
    if "xlsx" in payload:
        return parse_xlsx(payload)
    return parse_csv(payload.get("csv"))


def parse_xlsx(payload):
    """Bound the archive before streaming XLSX cells; reuse CSV cleaning."""
    encoded = payload.get("xlsx")
    if not isinstance(encoded, str):
        raise LabError("Excel cần dữ liệu XLSX hợp lệ.", 400)
    if len(encoded) > 4 * ((MAX_CSV_BYTES + 2) // 3):
        raise LabError("Tệp Excel vượt giới hạn 2 MiB.", 413)
    try:
        raw = base64.b64decode(encoded, validate=True)
        if len(raw) > MAX_CSV_BYTES:
            raise LabError("Tệp Excel vượt giới hạn 2 MiB.", 413)
        if not raw.startswith(b"PK\x03\x04"):
            raise ValueError()
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
            if len(entries) > 512 or sum(e.file_size for e in entries) > 32 * 1024 * 1024:
                raise LabError("Workbook Excel quá lớn sau giải nén (tối đa 32 MiB / 512 mục).", 413)
            if "xl/workbook.xml" not in archive.namelist():
                raise ValueError()
            for entry in entries:
                if entry.filename.endswith((".xml", ".rels")):
                    xml = archive.read(entry)
                    if b"<!DOCTYPE" in xml.upper() or b"<!ENTITY" in xml.upper():
                        raise ValueError()
                    # Check worksheet cells regardless of relationship target path.
                    if entry.filename.endswith(".xml"):
                        rows = 0
                        for _, element in ElementTree.iterparse(io.BytesIO(xml), events=("end",)):
                            tag = element.tag.rsplit("}", 1)[-1]
                            if tag == "row" and "spreadsheetml" in element.tag:
                                rows += 1
                                if rows > MAX_ROWS + 1 or int(element.get("r", "0")) > MAX_ROWS + 1:
                                    raise LabError("Sheet Excel vượt giới hạn 20.000 dòng dữ liệu.", 413)
                            elif tag == "c" and "spreadsheetml" in element.tag:
                                ref = re.fullmatch(r"([A-Z]+)(\d+)", element.get("r", ""))
                                if ref and int(ref[2]) > MAX_ROWS + 1:
                                    raise LabError("Sheet Excel vượt giới hạn 20.000 dòng dữ liệu.", 413)
                                column = 0
                                for letter in ref[1] if ref else "":
                                    column = column * 26 + ord(letter) - 64
                                if column > MAX_COLUMNS:
                                    raise LabError("Sheet Excel vượt giới hạn 100 cột.", 413)
                            element.clear()
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        try:
            sheets = workbook.sheetnames
            selected = payload.get("sheet", sheets[0] if sheets else None)
            if not isinstance(selected, str) or selected not in sheets:
                raise LabError("Sheet Excel đã chọn không tồn tại.")
            sheet = workbook[selected]
            # Read actual cells, even when worksheet dimension metadata is stale.
            # The archive scan above already bounds every sheet's cell coordinates.
            sheet.reset_dimensions()
            cells, width = [], 0
            for row in sheet.iter_rows():
                if len(cells) >= MAX_ROWS + 1 or len(row) > MAX_COLUMNS:
                    raise LabError("Sheet Excel vượt giới hạn 20.000 dòng / 100 cột.", 413)
                values = ["" if cell.value is None or cell.data_type == "e" else str(cell.value) for cell in row]
                while values and values[-1] == "":
                    values.pop()
                width = max(width, len(values))
                cells.append(values)
            # Match read_excel's rectangular table and omission of trailing blanks.
            while cells and not cells[-1]:
                cells.pop()
            output = io.StringIO()
            writer = csv.writer(output)
            writer.writerows(row + [""] * (width - len(row)) for row in cells)
        finally:
            workbook.close()
        # Same header/missing/blank/duplicate rules as CSV, then the same ML path.
        headers, rows, audit = parse_csv(output.getvalue(), check_size=False, delimiter=",")
        audit.update(file_type="xlsx", sheets=sheets, selected_sheet=selected)
        return headers, rows, audit
    except LabError:
        raise
    except Exception:
        raise LabError("Không đọc được Excel .xlsx hợp lệ; kiểm tra tệp và sheet đã chọn.", 422) from None


def leakage_warnings(headers, rows, target, features):
    t = headers.index(target)
    messages = []
    for name in features:
        i = headers.index(name)
        pairs = [(r[i], r[t]) for r in rows if r[i] is not None and r[t] is not None]
        if not pairs:
            continue
        mapping = {}
        for a, b in pairs:
            mapping.setdefault(a, set()).add(b)
        if all(a == b for a, b in pairs):
            messages.append({"column": name, "reason": "bản sao của target; cần loại trước khi chạy.", "severity": "hard"})
        elif len(pairs) >= 10 and len(mapping) <= min(30, len(pairs) // 5) and all(len(v) == 1 for v in mapping.values()) and len(set(b for _, b in pairs)) > 1:
            messages.append({"column": name, "reason": "có thể là bản mã hóa của target; quan hệ quan sát chưa chứng minh nguồn leakage.", "severity": "warning"})
        else:
            nums = [(number(a), number(b)) for a, b in pairs]
            nums = [(a, b) for a, b in nums if a is not None and b is not None]
            if len(nums) >= 10 and np.std([a for a, _ in nums]) > 0 and np.std([b for _, b in nums]) > 0:
                if abs(np.corrcoef(np.array(nums).T)[0, 1]) >= .995:
                    messages.append({"column": name, "reason": "tương quan gần tuyệt đối với target; có thể rò rỉ dữ liệu.", "severity": "warning"})
        if re.search(r"outcome|post_|after_|result|final_status", name, re.I):
            messages.append({"column": name, "reason": "tên gợi ý thông tin sau kết quả; tên cột chưa chứng minh leakage.", "severity": "warning"})
    return messages


def evaluation_plan(y, payload, headers, rows):
    n = len(y)
    task, strategy = payload["task"], payload.get("split", "random")
    tier = "normal" if n >= 100 else "small" if n >= 30 else "tiny" if n >= 10 else "extreme_tiny"
    default_folds = 5 if n >= 100 else 3
    indices, test = np.arange(n), np.array([], dtype=int)
    mode, notes = "cv_only", []
    dates = None
    if strategy == "temporal":
        col = payload.get("time_column")
        if col not in headers or col == payload.get("target"):
            raise LabError("Temporal split cần cột thời gian khác target.")
        dates = [date_key(r[headers.index(col)]) for r in rows]
        if any(d is None for d in dates):
            raise LabError("Cột thời gian phải có ISO date hợp lệ ở mọi dòng dùng để train.")
        indices = np.array(sorted(indices, key=lambda i: dates[i]))
    elif strategy != "random":
        raise LabError("Chiến lược chia dữ liệu không hợp lệ.")
    train = indices
    if n >= 30:
        if strategy == "temporal":
            cut = int(n * .75)
            while cut > 0 and dates[indices[cut - 1]] == dates[indices[cut]]:
                cut -= 1
            if cut >= 10 and n - cut >= 2:
                train, test = indices[:cut], indices[cut:]
        else:
            feasible = True
            if task == "classification":
                _, counts = np.unique(y, return_counts=True)
                feasible = min(counts) >= 3 and math.ceil(n * .25) >= len(counts) and math.floor(n * .75) >= len(counts)
            if feasible:
                train, test = train_test_split(indices, test_size=.25, random_state=42, stratify=y if task == "classification" else None)
        if task == "classification" and len(test):
            _, counts = np.unique(y[train], return_counts=True)
            if min(counts) < 2 or not set(y[test]).issubset(set(y[train])) or strategy == "random" and set(y[test]) != set(y):
                train, test = indices, np.array([], dtype=int)
        if len(test):
            mode = "holdout_cv"
    yt = y[train]
    folds = min(default_folds, len(train))
    if strategy == "temporal":
        # Split timestamp groups so equal timestamps never cross a CV boundary.
        groups = []
        for pos, row_index in enumerate(train):
            if not groups or dates[train[groups[-1][0]]] != dates[row_index]:
                groups.append([])
            groups[-1].append(pos)
        folds = min(default_folds, len(groups) - 1)
        splits = []
        if folds >= 2:
            for a, b in TimeSeriesSplit(n_splits=folds).split(groups):
                a = np.array([pos for g in a for pos in groups[g]])
                b = np.array([pos for g in b for pos in groups[g]])
                if task != "classification" or len(np.unique(yt[a])) >= 2 and set(yt[b]).issubset(set(yt[a])):
                    splits.append((a, b))
        cv_name = "TimeSeriesSplit"
        if len(splits) != max(0, folds):
            notes.append("Một số fold thời gian chưa đủ lớp để đánh giá an toàn và đã được bỏ qua.")
    elif task == "classification":
        _, counts = np.unique(yt, return_counts=True)
        folds = min(folds, int(min(counts)))
        splits = list(StratifiedKFold(folds, shuffle=True, random_state=42).split(train, yt)) if folds >= 2 else []
        cv_name = "StratifiedKFold"
    else:
        folds = min(folds, max(2, len(train) // 2))
        splits = list(KFold(folds, shuffle=True, random_state=42).split(train)) if len(train) >= 4 else []
        cv_name = "KFold"
    if len(splits) < 2:
        mode, train, test, splits, cv_name = "exploratory", indices, np.array([], dtype=int), [], "None"
        notes.append("Exploratory / không có đánh giá độc lập. Không đủ mẫu hoặc mẫu mỗi lớp để tạo CV an toàn; chỉ báo cáo kết quả trên dữ liệu đã dùng để fit.")
    elif not len(test):
        notes.append("Không tạo test set riêng vì dữ liệu quá nhỏ hoặc quá ít mẫu mỗi lớp. Hệ thống đã chuyển sang Adaptive Cross-Validation; không có held-out test metric.")
    label = "Exploratory Only" if mode == "exploratory" else "Tiny Data Adaptive CV" if tier in ("tiny", "extreme_tiny") else "Normal Holdout + CV" if tier == "normal" and mode == "holdout_cv" else f"Small Data {len(splits)}-Fold CV" if mode == "holdout_cv" else "CV-only — Adaptive CV"
    evaluation = dict(mode=mode, label=label, tier=tier, usable_rows=n, class_count=len(np.unique(y)) if task == "classification" else None,
                      folds=len(splits), has_holdout=bool(len(test)), independent_evaluation=mode != "exploratory")
    return train, test, splits, cv_name, evaluation, notes


def prepare(payload):
    headers, rows, audit = load_data(payload)
    summary = profile(headers, rows, audit)
    target = payload.get("target")
    if target not in headers:
        raise LabError("Hãy chọn cột mục tiêu hợp lệ.")
    info = summary["columns"][headers.index(target)]
    if info["unique"] < 2:
        raise LabError("Target toàn thiếu hoặc chỉ có một giá trị; hãy chọn cột khác.")
    task = payload.get("task")
    if task not in ("classification", "regression"):
        raise LabError("Bạn cần xác nhận Classification hoặc Regression.")
    if payload.get("date_features", "drop") not in ("drop", "extract"):
        raise LabError("Cách xử lý cột ngày không hợp lệ.")
    features = payload.get("features")
    if not isinstance(features, list) or not features or any(f not in headers or f == target for f in features) or len(set(features)) != len(features):
        raise LabError("Chọn ít nhất một feature; target không được nằm trong features.")
    t = headers.index(target)
    target_format = column_numeric_format([r[t] for r in rows])
    if task == "regression" and info["date"]:
        raise LabError("Target là ngày/giờ, không phải giá trị số. Hãy tạo duration / số ngày giữa hai mốc có ý nghĩa trước khi dùng Regression.")
    parsed_target = [number(r[t], target_format) for r in rows]
    if task == "regression":
        audit["numeric_target"] = {"total_rows": len(rows), "nonblank_target_rows": sum(r[t] is not None for r in rows),
            "valid_numeric_target_rows": sum(v is not None for v in parsed_target), "missing_target_rows": sum(r[t] is None for r in rows),
            "invalid_numeric_target_rows": sum(r[t] is not None and v is None for r, v in zip(rows, parsed_target)),
            "unique_numeric_target_values": len({v for v in parsed_target if v is not None})}
    usable = [r for r, v in zip(rows, parsed_target) if r[t] is not None and (task != "regression" or v is not None)]
    audit["target_rows_removed"] = len(rows) - len(usable)
    if len(usable) < 2:
        raise LabError("Số mẫu hợp lệ quá ít để fit mô hình; cần ít nhất 2 dòng với target khác nhau. " + ("Numeric target diagnostics: " + json.dumps(audit["numeric_target"]) if task == "regression" else ""))
    if task == "regression" and not any(number(r[t]) is not None for r in rows):
        raise LabError("Regression yêu cầu target là số.")
    y = np.array([r[t] for r in usable]) if task == "classification" else np.array([number(r[t], target_format) for r in usable])
    if len(np.unique(y)) < 2:
        raise LabError("Target sau làm sạch chỉ có một giá trị.")
    if task == "classification" and len(np.unique(y)) > MAX_CLASSES:
        if info["date"]:
            raise LabError("Target có vẻ là cột ngày tháng và có quá nhiều giá trị riêng biệt để dùng trực tiếp cho Classification. Hãy chọn target nhãn như outcome/gender hoặc chuyển ngày thành khoảng thời gian/số ngày giữa hai mốc để dùng Regression.")
        if info["id_like"]:
            raise LabError("Target có quá nhiều lớp và có vẻ giống cột định danh. Hãy chọn target nhãn có ý nghĩa.")
        raise LabError("Classification hỗ trợ tối đa 30 lớp; kiểm tra target hoặc chọn Regression.")
    strategy = payload.get("split", "random")
    train, test, splits, cv_name, evaluation, notes = evaluation_plan(y, payload, headers, usable)
    leakage = leakage_warnings(headers, [usable[j] for j in train], target, features)
    copies = [m["column"] for m in leakage if m["severity"] == "hard"]
    excluded = {m["column"] for m in leakage if m["severity"] == "hard" or not payload.get("ack_leakage")}
    features = [f for f in features if f not in excluded]
    messages = [f"{m['column']}: {m['reason']} " + ("Đã loại và tiếp tục." if m['column'] in excluded else "Đã xác nhận nguồn feature.") for m in leakage]
    if not features:
        raise LabError("Không còn feature sau khi loại các cột có nguy cơ leakage. Hãy chọn feature khác hoặc xác nhận nguồn các heuristic warning; bản sao target luôn phải loại.")
    if copies:
        messages.append("Đã loại feature là bản sao của target để tránh leakage: " + ", ".join(copies))
    messages.extend(audit.get("warnings", []))
    if audit["target_rows_removed"]:
        messages.append(f"Đã loại {audit['target_rows_removed']} dòng vì target bị thiếu hoặc không phải số hợp lệ cho Regression. Không impute target.")
    if info["id_like"] or info["date"]:
        messages.append("Target có dạng ID/ngày: đã giữ lựa chọn của bạn; kiểm tra ý nghĩa bài toán trước khi dùng kết quả.")
    messages.extend(notes)
    # Infer numerical/categorical types only from TRAIN, never from test.
    cols, names, numeric, categorical, selected = [], [], [], [], []
    for name in features:
        i = headers.index(name)
        values = [usable[j][i] for j in train if usable[j][i] is not None]
        numeric_format = column_numeric_format(values)
        if not values:
            messages.append(f"Đã bỏ {name}: toàn thiếu trong train.")
            continue
        if sum(bool(DATE_LIKE.match(v) or TIME_LIKE.match(v)) for v in values) / len(values) >= .9:
            if payload.get("date_features", "drop") == "drop":
                messages.append(f"Đã bỏ cột ngày {name} theo lựa chọn của bạn.")
                continue
            if any(date(v) is None for v in values):
                messages.append(f"Đã bỏ {name}: ngày mơ hồ hoặc không hợp lệ; chuẩn hóa ISO để trích xuất ngày.")
                continue
            for part in ("year", "month", "day", "day_of_week"):
                arr = []
                for row in usable:
                    d = date(row[i])
                    arr.append(np.nan if d is None else d.weekday() if part == "day_of_week" else getattr(d, part))
                numeric.append(len(cols)); cols.append(arr); names.append(f"{name}__{part}")
        elif (numeric_format or not any(ambiguous_number(v) for v in values)) and sum(number(v, numeric_format) is not None for v in values) / len(values) >= .9:
            numeric.append(len(cols)); cols.append([number(r[i], numeric_format) if number(r[i], numeric_format) is not None else np.nan for r in usable]); names.append(name)
        else:
            categorical.append(len(cols)); cols.append([r[i] if r[i] is not None else np.nan for r in usable]); names.append(name)
            if len(set(values)) > 50:
                messages.append(f"{name}: nhiều category; OHE gộp nhóm hiếm (tối đa 32 category/cột, học trên train mỗi fold).")
        selected.append(name)
    if not cols:
        raise LabError("Không còn feature sau xử lý ngày; hãy chọn thêm feature.")
    expanded = len(numeric) + sum(min(32, len({cols[i][j] for j in train if isinstance(cols[i][j], str)})) for i in categorical)
    feature_guard = {"estimated_before": expanded, "excluded_columns": [], "reason": "OHE bounded at 32 categories per column, fit within each training fold"}
    max_categories = 32
    if expanded > MAX_EXPANDED or expanded * len(usable) > MAX_CELLS:
        budget = min(MAX_EXPANDED, MAX_CELLS // len(usable))
        max_categories = min(32, (budget - len(numeric)) // max(1, len(categorical)))
        if max_categories < 2:
            # Keep numerical columns and as many bounded categorical columns as fit.
            keep = numeric + categorical[:max(0, (budget - len(numeric)) // 2)]
            dropped = [names[i] for i in range(len(cols)) if i not in keep]
            feature_guard["excluded_columns"] = dropped
            cols, names = [cols[i] for i in keep], [names[i] for i in keep]
            numeric, categorical = list(range(len(numeric))), list(range(len(numeric), len(keep)))
            selected = names.copy()
            max_categories = 2
            messages.append("Đã bỏ cột categorical vượt ngân sách feature: " + ", ".join(dropped))
        if not cols:
            raise LabError("Không còn feature trong giới hạn tài nguyên serverless; hãy chọn ít cột hơn.")
        expanded = len(numeric) + len(categorical) * max_categories
        messages.append(f"Đã gộp category hiếm, tối đa {max_categories} nhóm/cột để giữ dưới {MAX_EXPANDED} feature / 3 triệu ô.")
    X = np.array(cols, dtype=object).T
    feature_guard["estimated_after"] = expanded
    if expanded != feature_guard["estimated_before"]:
        feature_guard["reason"] = "Encoded feature / cell budget; reduced rare category groups and excluded overflow columns"
    yt = y[train]
    folds = len(splits)
    if len(usable) < 10:
        messages.append("Dữ liệu rất ít. Kết quả chỉ mang tính thử nghiệm, không nên dùng để lựa chọn mô hình thực tế.")
    elif len(usable) < 30:
        messages.append("Dữ liệu rất nhỏ: Adaptive CV có độ bất định cao; metric có thể biến động rất mạnh.")
    elif len(usable) < 100:
        messages.append("Dữ liệu khá nhỏ (dưới 100 dòng hợp lệ). Kết quả Cross Validation và Test có thể biến động mạnh.")
    if task == "classification":
        labels, counts = np.unique(yt, return_counts=True)
        if min(counts) / sum(counts) < .15:
            messages.append("Mất cân bằng lớp: F1 Macro ưu tiên công bằng giữa lớp; model hỗ trợ dùng class_weight=balanced.")
    summary.pop("preview", None)
    return dict(X=X, y=y, train=train, test=test, splits=splits, task=task, numeric=numeric, categorical=categorical,
                names=names, summary=summary, warnings=messages, selected_features=selected, split=strategy,
                expanded_estimate=expanded, feature_guard=feature_guard, leakage=leakage, folds=folds, max_categories=max_categories, cv_name=cv_name, evaluation=evaluation)


def preprocessor(data, scale):
    transformers = []
    if data["numeric"]:
        steps = [("imputer", SimpleImputer(strategy="median", keep_empty_features=True))]
        if scale:
            steps.append(("scaler", StandardScaler()))
        transformers.append(("numeric", Pipeline(steps), data["numeric"]))
    if data["categorical"]:
        transformers.append(("categorical", Pipeline([
            ("imputer", SimpleImputer(strategy="most_frequent", keep_empty_features=True)),
            ("onehot", OneHotEncoder(handle_unknown="ignore", sparse_output=False, min_frequency=2, max_categories=data.get("max_categories", 32)))
        ]), data["categorical"]))
    return ColumnTransformer(transformers, verbose_feature_names_out=False)


def models(task):
    if task == "classification":
        return [("Logistic Regression", LogisticRegression(max_iter=600, class_weight="balanced", random_state=42), True),
                ("K-Nearest Neighbors", KNeighborsClassifier(n_neighbors=5, n_jobs=1), True),
                ("Support Vector Machine", SVC(class_weight="balanced", random_state=42), True),
                ("Decision Tree", DecisionTreeClassifier(max_depth=8, min_samples_leaf=2, class_weight="balanced", random_state=42), False),
                ("Random Forest", RandomForestClassifier(n_estimators=60, max_depth=12, min_samples_leaf=2, class_weight="balanced", n_jobs=1, random_state=42), False),
                ("Gaussian Naive Bayes", GaussianNB(), False),
                ("Gradient Boosting", GradientBoostingClassifier(n_estimators=60, max_depth=2, random_state=42), False)]
    return [("Linear Regression", LinearRegression(), True), ("Ridge", Ridge(alpha=1), True),
            ("Lasso", Lasso(alpha=.01, max_iter=2000), True), ("KNN Regressor", KNeighborsRegressor(n_neighbors=5, n_jobs=1), True),
            ("SVR", SVR(), True), ("Decision Tree Regressor", DecisionTreeRegressor(max_depth=8, min_samples_leaf=2, random_state=42), False),
            ("Random Forest Regressor", RandomForestRegressor(n_estimators=60, max_depth=12, min_samples_leaf=2, n_jobs=1, random_state=42), False),
            ("Gradient Boosting Regressor", GradientBoostingRegressor(n_estimators=60, max_depth=2, random_state=42), False)]


def primary(task, y, pred):
    return float(metrics.f1_score(y, pred, average="macro", zero_division=0)) if task == "classification" else float(np.sqrt(metrics.mean_squared_error(y, pred)))


def evaluate(pipe, X, y, task):
    started = time.perf_counter()
    pred = pipe.predict(X)
    elapsed = (time.perf_counter() - started) * 1000
    if task == "regression":
        mse = metrics.mean_squared_error(y, pred)
        result = dict(mae=float(metrics.mean_absolute_error(y, pred)), mse=float(mse), rmse=float(np.sqrt(mse)), r2=float(metrics.r2_score(y, pred)))
        if np.min(np.abs(y)) > 1e-8:
            result["mape"] = float(metrics.mean_absolute_percentage_error(y, pred))
        return result, pred, elapsed, None
    result = {"accuracy": float(metrics.accuracy_score(y, pred)),
              "precision": float(metrics.precision_score(y, pred, average="macro", zero_division=0)),
              "recall": float(metrics.recall_score(y, pred, average="macro", zero_division=0)),
              "f1": float(metrics.f1_score(y, pred, average="macro", zero_division=0)),
              "f1_weighted": float(metrics.f1_score(y, pred, average="weighted", zero_division=0)),
              "roc_auc": None, "pr_auc": None, "log_loss": None}
    curves = None
    labels = pipe.classes_
    probs = pipe.predict_proba(X) if hasattr(pipe, "predict_proba") else None
    scores = pipe.decision_function(X) if probs is None and hasattr(pipe, "decision_function") else None
    try:
        if len(labels) == 2 and len(np.unique(y)) == 2 and (probs is not None or scores is not None):
            positive = (y == labels[1]).astype(int)
            score = probs[:, 1] if probs is not None else scores
            result["roc_auc"] = float(metrics.roc_auc_score(positive, score))
            precision, recall, _ = metrics.precision_recall_curve(positive, score)
            result["pr_auc"] = float(metrics.auc(recall, precision))
            fpr, tpr, _ = metrics.roc_curve(positive, score)
            curves = {"positive_class": str(labels[1]), "roc": list(map(list, zip(fpr.tolist(), tpr.tolist()))),
                      "pr": list(map(list, zip(recall.tolist(), precision.tolist())))}
        elif probs is not None and set(y) == set(labels):
            result["roc_auc"] = float(metrics.roc_auc_score(y, probs, labels=labels, multi_class="ovr", average="macro"))
        if probs is not None:
            result["log_loss"] = float(metrics.log_loss(y, probs, labels=labels))
    except ValueError:
        pass  # Unavailable metric is null, never a made-up value.
    return result, pred, elapsed, curves


def rank_models(records, task):
    # Only train CV fields are read. Test metrics cannot change this ordering.
    return sorted(records, key=lambda r: ((-1 if task == "classification" else 1) * r["cv_mean"], r["cv_std"], r["cv_fit_ms"], r["name"]))


def fit_with_budget(pipe, X, y, deadline):
    remaining = deadline - time.perf_counter()
    if remaining <= 0:
        raise TimeoutError()
    # Linux worker runs on its main thread. A POSIX alarm interrupts an
    # over-budget fit; the existing subprocess deadline remains the hard guard.
    if hasattr(signal, "setitimer") and __import__("threading").current_thread() is __import__("threading").main_thread():
        previous = signal.getsignal(signal.SIGALRM)
        def expired(signum, frame):
            raise TimeoutError()
        signal.signal(signal.SIGALRM, expired)
        signal.setitimer(signal.ITIMER_REAL, min(15, remaining))
        try:
            return pipe.fit(X, y)
        finally:
            signal.setitimer(signal.ITIMER_REAL, 0)
            signal.signal(signal.SIGALRM, previous)
    return pipe.fit(X, y)


def benchmark(payload):
    started = time.perf_counter()
    logger.info("ml_compare stage=prepare")
    data = prepare(payload)
    X, y, train, test = (data[k] for k in ("X", "y", "train", "test"))
    task, skipped, trained, records = data["task"], [], {}, []
    exploratory = data["evaluation"]["mode"] == "exploratory"
    safe_train_size = min([len(a) for a, _ in data["splits"]] or [len(train)])
    with threadpool_limits(limits=1), warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        for name, estimator, scale in models(task):
            reason = None
            if len(y) > 4000 and name in ("Support Vector Machine", "SVR", "K-Nearest Neighbors", "KNN Regressor"):
                reason = "Dataset quá lớn để benchmark model này an toàn trên serverless."
            elif time.perf_counter() - started > 35:
                reason = "Không chạy do giới hạn thời gian; so sánh chỉ áp dụng các model hoàn tất."
            if reason:
                skipped.append({"name": name, "status": "SKIPPED_TIME_BUDGET", "reason": reason}); continue
            if isinstance(estimator, (KNeighborsClassifier, KNeighborsRegressor)):
                neighbors = min(5, max(1, safe_train_size - 1))
                estimator.set_params(n_neighbors=neighbors)
                if neighbors != 5:
                    data["warnings"].append(f"{name}: đã điều chỉnh n_neighbors={neighbors} theo số mẫu train nhỏ nhất trong CV.")
            pipe = Pipeline([("preprocess", preprocessor(data, scale)), ("model", estimator)])
            try:
                logger.info("ml_compare stage=cv model=%s", type(estimator).__name__)
                values, cv_fit = [], 0
                for a, b in data["splits"]:
                    p = clone(pipe)
                    clock = time.perf_counter()
                    fit_with_budget(p, X[train][a], y[train][a], started + 45)
                    cv_fit += (time.perf_counter() - clock) * 1000
                    values.append(primary(task, y[train][b], p.predict(X[train][b])))
                    if time.perf_counter() - started > 44:
                        raise TimeoutError()
                clock = time.perf_counter()
                logger.info("ml_compare stage=fit_final model=%s", type(estimator).__name__)
                fit_with_budget(pipe, X[train], y[train], started + 45)
                fit_ms = (time.perf_counter() - clock) * 1000
                record = {"name": name, "status": "SUCCESS", "algorithm_type": type(estimator).__name__, "cv_mean": float(np.mean(values)) if values else None,
                          "cv_std": float(np.std(values)) if values else None, "cv_scores": values, "cv_fit_ms": cv_fit, "fit_time_ms": fit_ms,
                          "exploratory_score": primary(task, y[train], pipe.predict(X[train])) if exploratory else None,
                          "hyperparameters": estimator.get_params(), "preprocessing": {"numeric": "median + StandardScaler" if scale else "median", "categorical": f"most_frequent + OneHotEncoder (max {data['max_categories']}/cột)"}}
                records.append(record); trained[name] = pipe
            except TimeoutError:
                skipped.append({"name": name, "status": "SKIPPED_TIME_BUDGET", "reason": "Đã dừng do ngân sách thời gian."})
            except Exception:
                logger.exception("ML comparison model failed algorithm=%s", type(estimator).__name__)
                skipped.append({"name": name, "status": "FAILED_SAFE", "reason": "Model không tương thích hoặc không hội tụ trên cấu hình dữ liệu này; các model còn lại vẫn tiếp tục."})
        if not records:
            raise LabError("Không có model hoàn tất. Giảm số dòng/feature hoặc kiểm tra dữ liệu.")
        records = sorted(records, key=lambda r: ((-1 if task == "classification" else 1) * r["exploratory_score"], r["name"])) if exploratory else rank_models(records, task)
        recommended = None if exploratory else records[0]["name"]  # Frozen BEFORE any held-out test evaluation.
        for rank, record in enumerate(records, 1):
            pipe = trained[record["name"]]
            # Missing independent metrics are genuinely absent, not train scores.
            record.update(rank=rank, train={}, test=None, predict_time_ms=None, generalization_gap=None,
                          overfit_risk="N/A", importance=None, diagnostics_source="heldout_test" if len(test) else "none",
                          model_size_bytes=None)
            try:
                record["model_size_bytes"] = len(pickle.dumps(pipe, protocol=5))
                add_diagnostics(record, pipe, data)
            except Exception:
                logger.exception("ML comparison diagnostics failed algorithm=%s", record["algorithm_type"])
                record["diagnostics_status"] = "FAILED_SAFE"
                data["warnings"].append(f"{record['name']}: không tính được đầy đủ diagnostics; điểm CV và thứ hạng vẫn giữ nguyên.")
        if any("ConvergenceWarning" in type(w.message).__name__ for w in caught):
            data["warnings"].append("Một số model đạt giới hạn vòng lặp; cân nhắc kết quả như baseline, không phải tuning tối ưu.")
    best = records[0]
    quality = {"level": "normal", "message": ""}
    if exploratory:
        quality = {"level": "caution", "message": "Exploratory / không có đánh giá độc lập. Điểm trên dữ liệu fit không chứng minh khả năng dự đoán; không dùng để lựa chọn mô hình thực tế."}
    elif not len(test):
        quality = {"level": "caution", "message": "Chỉ có Adaptive CV, không có test set riêng. Mẫu ít nên kết quả có độ bất định cao; cần thêm dữ liệu để kiểm chứng độc lập."}
    elif best.get("test"):
        if task == "regression":
            score = best["test"]["r2"]
            if score is not None and score < .1:
                quality = {"level": "low" if score < 0 else "caution", "message": "Mô hình tốt nhất theo CV nhưng khả năng tổng quát hóa trên test còn yếu. Dataset có thể quá nhỏ hoặc thiếu feature hữu ích. Không nên sử dụng cho quyết định thực tế khi chưa cải thiện dữ liệu."}
        else:
            labels, counts = np.unique(y[train], return_counts=True)
            baseline = float(metrics.f1_score(y[test], np.repeat(labels[np.argmax(counts)], len(test)), average="macro", zero_division=0))
            score = best["test"]["f1"]
            if score < .4 or score <= baseline + .05:
                quality = {"level": "low", "message": "F1 Macro trên test thấp hoặc gần baseline lớp phổ biến nhất. Dữ liệu hiện tại chưa chứng minh khả năng dự đoán đáng tin cậy; cần cải thiện trước khi dùng thực tế."}
            quality["baseline_test_f1"] = baseline
    if skipped:
        data["warnings"].append(f"{len(skipped)} mô hình không thể chạy với cấu trúc dữ liệu hoặc ngân sách hiện tại và đã được bỏ qua.")
    return {"task": task, "recommended": recommended, "best_exploratory_model": best["name"] if exploratory else None,
            "evaluation": data["evaluation"], "leakage": data["leakage"], "quality": quality, "primary_metric": "F1 Macro" if task == "classification" else "RMSE",
            "selection_source": "exploratory_train" if exploratory else "train_cv" if len(test) else "cv_only", "models": records, "skipped": skipped, "summary": data["summary"], "warnings": data["warnings"],
            "methodology": {"split": data["split"], "train_rows": len(train), "test_rows": len(test), "folds": data["folds"],
                            "evaluation_mode": data["evaluation"]["mode"], "cv": data["cv_name"],
                            "seed": 42, "selected_features": data["selected_features"], "expanded_estimate": data["expanded_estimate"],
                            "feature_guard": data["feature_guard"],
                            "preprocessing_fit": "train only, refit within every CV fold" if not exploratory else "fit data only; no independent evaluation", "test_used_for_selection": False},
            "duration_ms": (time.perf_counter() - started) * 1000, "raw_file_persisted": False}


def add_diagnostics(record, pipe, data):
    X, y, train, test, task = (data[k] for k in ("X", "y", "train", "test", "task"))
    train_metrics, _, _, _ = evaluate(pipe, X[train], y[train], task)
    record["train"] = train_metrics
    if not len(test):
        return
    test_metrics, pred, predict_ms, curves = evaluate(pipe, X[test], y[test], task)
    gap = train_metrics["f1"] - test_metrics["f1"] if task == "classification" else train_metrics["r2"] - test_metrics["r2"]
    record.update(train=train_metrics, test=test_metrics, predict_time_ms=predict_ms,
                  model_size_bytes=len(pickle.dumps(pipe, protocol=5)), generalization_gap=float(gap),
                  overfit_risk="High" if gap > .15 else "Medium" if gap > .05 else "Low")
    if task == "classification":
        labels = pipe.classes_
        cm = metrics.confusion_matrix(y[test], pred, labels=labels)
        confused = sorted([(int(cm[i, j]), str(labels[i]), str(labels[j])) for i in range(len(labels)) for j in range(len(labels)) if i != j], reverse=True)[:3]
        record.update(confusion={"labels": labels.tolist(), "matrix": cm.tolist()},
                      classification_report=metrics.classification_report(y[test], pred, labels=labels, output_dict=True, zero_division=0),
                      curves=curves, most_confused=[{"count": n, "actual": a, "predicted": b} for n, a, b in confused if n])
    else:
        sampled = np.linspace(0, len(test) - 1, min(250, len(test)), dtype=int)
        residual = y[test] - pred
        record.update(scatter=[[float(y[test][i]), float(pred[i])] for i in sampled],
                      residuals=[[float(pred[i]), float(residual[i])] for i in sampled],
                      error_analysis={"mean_residual": float(np.mean(residual)), "largest_absolute_residual": float(np.max(np.abs(residual))), "convention": "actual - predicted"})
    estimator = pipe.named_steps["model"]
    weights = getattr(estimator, "feature_importances_", None)
    if weights is None and hasattr(estimator, "coef_"):
        weights = np.abs(estimator.coef_)
        if weights.ndim > 1:
            weights = np.mean(weights, axis=0)
    record["importance"] = None
    if weights is not None:
        feature_names = pipe.named_steps["preprocess"].get_feature_names_out(data["names"])
        record["importance"] = sorted([{"feature": str(n), "value": float(w)} for n, w in zip(feature_names, weights)], key=lambda x: -x["value"])[:15]


def sanitize(value):
    if isinstance(value, dict):
        return {str(k): sanitize(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [sanitize(v) for v in value]
    if isinstance(value, (float, np.floating)):
        return float(value) if math.isfinite(value) else None
    if isinstance(value, np.integer):
        return int(value)
    return value


def _worker_environment():
    env = os.environ.copy()
    paths = [item for item in sys.path if isinstance(item, str) and item]
    paths.extend(item for item in env.get("PYTHONPATH", "").split(os.pathsep) if item)
    env["PYTHONPATH"] = os.pathsep.join(dict.fromkeys(paths))
    return env


def dispatch(payload, isolated=True):
    if not isinstance(payload, dict):
        raise LabError("Request JSON không hợp lệ.", 400)
    action = payload.get("action")
    if action == "inspect":
        headers, rows, audit = load_data(payload)
        result = profile(headers, rows, audit)
        if payload.get("target") in headers:
            result["leakage"] = leakage_warnings(headers, rows, payload["target"], [c for c in headers if c != payload["target"]])
        return result
    if action != "benchmark":
        raise LabError("Action không hợp lệ.", 400)
    if not isolated:
        return benchmark(payload)
    # No request data is explicitly logged or written to argv/files. Worker stderr
    # reaches server logs, including import failures before its exception handler.
    # Killable process bounds
    # even a single estimator stuck inside native code. No warm cache of user data.
    try:
        logger.info("ml_compare stage=worker_start")
        completed = subprocess.run([sys.executable, "-B", str(Path(__file__).resolve()), "--worker"],
                                   input=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
                                   stdout=subprocess.PIPE, stderr=None, timeout=55,
                                   env=_worker_environment(),
                                   creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0)
        if completed.returncode != 0:
            logger.error("ML comparison worker exited with code=%s; see worker stderr", completed.returncode)
            raise LabError("Huấn luyện không hoàn tất; hãy kiểm tra dữ liệu hoặc giảm kích thước.", 500)
        result = json.loads(completed.stdout)
        if "error" in result:
            raise LabError(result["error"], result.get("status", 500))
        return result
    except subprocess.TimeoutExpired:
        raise LabError("Đã dừng huấn luyện sau 55 giây. Hãy giảm số dòng hoặc feature.", 504) from None


class handler(BaseHTTPRequestHandler):
    def log_message(self, *_args):
        pass

    def respond(self, status, value):
        body = json.dumps(sanitize(value), ensure_ascii=False, allow_nan=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self):
        self.respond(200, {"available": True, "backend": "scikit-learn", "max_csv_bytes": MAX_CSV_BYTES,
                           "max_rows": MAX_ROWS, "max_columns": MAX_COLUMNS, "raw_file_persisted": False})

    def do_POST(self):
        try:
            if self.headers.get("Content-Type", "").split(";", 1)[0].strip().lower() != "application/json":
                raise LabError("Request cần Content-Type application/json.", 400)
            try:
                size = int(self.headers.get("Content-Length", "0"))
            except ValueError:
                raise LabError("Content-Length không hợp lệ.", 400) from None
            if size > MAX_REQUEST_BYTES:
                raise LabError("Request vượt giới hạn 4 MiB; hãy giảm kích thước tệp.", 413)
            if size <= 0:
                raise LabError("Request rỗng.", 400)
            payload = json.loads(self.rfile.read(size).decode("utf-8"))
            self.respond(200, dispatch(payload))
        except LabError as error:
            self.respond(error.status, {"error": str(error)})
        except (UnicodeError, json.JSONDecodeError):
            self.respond(400, {"error": "JSON/UTF-8 không hợp lệ."})
        except Exception:
            logger.exception("ML comparison request failed")
            self.respond(500, {"error": "Dịch vụ phân tích gặp lỗi; hãy thử dữ liệu nhỏ hơn."})

    def do_PUT(self):
        self.respond(405, {"error": "Chỉ hỗ trợ GET và POST."})


def run_worker():
    try:
        output = benchmark(json.loads(sys.stdin.buffer.read().decode("utf-8")))
    except LabError as error:
        output = {"error": str(error), "status": error.status}
    except Exception:
        logger.exception("ML comparison benchmark failed")
        output = {"error": "Không thể huấn luyện với dữ liệu này.", "status": 500}
    sys.stdout.buffer.write(json.dumps(sanitize(output), ensure_ascii=False, allow_nan=False).encode("utf-8"))


if __name__ == "__main__":
    if "--worker" in sys.argv:
        run_worker()
    elif "--serve" in sys.argv:
        class Local(handler, SimpleHTTPRequestHandler):
            def do_GET(self):
                if self.path.split("?")[0] == "/api/ml-compare":
                    return handler.do_GET(self)
                return SimpleHTTPRequestHandler.do_GET(self)

            def do_POST(self):
                if self.path != "/api/ml-compare":
                    return self.respond(404, {"error": "Không tìm thấy endpoint."})
                return handler.do_POST(self)
        os.chdir(Path(__file__).resolve().parents[1])
        port = int(os.environ.get("ML_COMPARE_PORT", "8000"))
        print(f"ML Comparison local server: http://127.0.0.1:{port}/ml-model-comparison.html", flush=True)
        ThreadingHTTPServer(("127.0.0.1", port), Local).serve_forever()
