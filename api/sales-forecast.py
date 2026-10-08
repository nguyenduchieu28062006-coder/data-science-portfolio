"""Stateless sales forecasting; no pandas, storage, random split or future targets.

Local: python -B api/sales-forecast.py --serve (port 8001).
Vercel: handler, with a killable 55-second worker and a 45-second soft budget.
"""
import base64
import calendar
import csv
import io
import json
import logging
import math
import os
import re
import subprocess
import sys
import time
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta, timezone
from http.server import BaseHTTPRequestHandler, SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

for _key in ("OMP_NUM_THREADS", "OPENBLAS_NUM_THREADS", "MKL_NUM_THREADS"):
    os.environ[_key] = "1"
import numpy as np
from sklearn.linear_model import LinearRegression, Ridge
from sklearn.pipeline import make_pipeline
from sklearn.preprocessing import StandardScaler
from sklearn.ensemble import RandomForestRegressor, GradientBoostingRegressor
try:
    from sklearn.ensemble import HistGradientBoostingRegressor
except ImportError:
    HistGradientBoostingRegressor = None
from threadpoolctl import threadpool_limits

MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_REQUEST_BYTES = 4 * 1024 * 1024
MAX_ROWS, MAX_COLUMNS, MAX_CELLS = 50000, 60, 1500000
MISSING = {"", "na", "n/a", "null", "none", "nan", "-", "--"}
MODELS = ["Naive Forecast", "Seasonal Naive", "Linear Regression", "Ridge Regression",
          "Random Forest", "Gradient Boosting", "Hist Gradient Boosting"]
PERIOD = {"hourly": 24, "daily": 7, "weekly": 52, "monthly": 12, "quarterly": 4}
logger = logging.getLogger("sales_forecast")


class ForecastError(Exception):
    def __init__(self, message, status=422):
        super().__init__(message)
        self.status = status


def token(value):
    return "" if value is None else str(value).strip()


def numeric(value, convention="auto"):
    text = token(value)
    if text.lower() in MISSING:
        return None
    if re.fullmatch(r"[+-]?\d{1,3}(?:,\d{3})+\.\d+", text):
        text = text.replace(",", "")
    elif re.fullmatch(r"[+-]?\d{1,3}(?:\.\d{3})+,\d+", text):
        text = text.replace(".", "").replace(",", ".")
    elif convention == "decimal_comma":
        if "." in text and "," not in text:
            if not re.fullmatch(r"[+-]?\d{1,3}(?:\.\d{3})+", text):
                return None
        text = text.replace(".", "").replace(",", ".")
    elif convention == "decimal_dot":
        if "," in text and not re.fullmatch(r"[+-]?\d{1,3}(?:,\d{3})+(?:\.\d+)?", text):
            return None
        text = text.replace(",", "")
    elif "," in text:
        if re.fullmatch(r"[+-]?\d{1,3},\d{3}", text):
            return None
        text = text.replace(",", ".")
    if not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", text):
        return None
    result = float(text)
    return result if math.isfinite(result) and abs(result) <= 1e12 else None


def date_order(values, requested="auto"):
    if requested not in ("auto", "day_first", "month_first"):
        raise ForecastError("Cách đọc ngày không hợp lệ.")
    if requested != "auto":
        return requested, False
    evidence = set()
    ambiguous = False
    for value in values:
        match = re.fullmatch(r"(\d{1,2})[-/](\d{1,2})[-/](\d{4})(?:[ T].*)?", token(value))
        if match:
            a, b = int(match[1]), int(match[2])
            if a > 12 and b <= 12:
                evidence.add("day_first")
            elif b > 12 and a <= 12:
                evidence.add("month_first")
            elif a <= 12 and b <= 12 and a != b:
                ambiguous = True
    if len(evidence) > 1:
        raise ForecastError("Ngày chứa cả DD/MM và MM/DD. Hãy chuẩn hóa sang YYYY-MM-DD.")
    return next(iter(evidence), "auto"), ambiguous and not evidence


def parse_date(value, order="auto", excel=False):
    if isinstance(value, datetime):
        result = value
    elif isinstance(value, date):
        result = datetime.combine(value, datetime.min.time())
    elif excel and isinstance(value, (float, int)) and 1 <= value <= 100000:
        from openpyxl.utils.datetime import from_excel
        result = from_excel(value)
        if not isinstance(result, datetime):
            return None
    else:
        text = token(value)
        try:
            if re.fullmatch(r"\d{4}[-/]\d{1,2}(?:[-/]\d{1,2})?(?:[ T].*)?", text):
                if re.fullmatch(r"\d{4}[-/]\d{1,2}", text):
                    text += "-01"
                result = datetime.fromisoformat(text.replace("/", "-").replace("Z", "+00:00"))
            else:
                match = re.fullmatch(r"(\d{1,2})[-/](\d{1,2})[-/](\d{4})(?:[ T](\d{2}:\d{2}(?::\d{2})?))?", text)
                if not match:
                    return None
                a, b, year = map(int, match.group(1, 2, 3))
                if order == "auto":
                    if a > 12:
                        order = "day_first"
                    elif b > 12:
                        order = "month_first"
                    elif a != b:
                        return None
                result = datetime(year, b, a) if order != "month_first" else datetime(year, a, b)
                if match[4]:
                    parts = list(map(int, match[4].split(":")))
                    result = result.replace(hour=parts[0], minute=parts[1], second=parts[2] if len(parts) == 3 else 0)
        except (ValueError, OverflowError):
            return None
    # Offset timestamps consistently use UTC; naive dates retain their clock time.
    return result.astimezone(timezone.utc).replace(tzinfo=None) if result.tzinfo else result


def sample_table():
    rng = np.random.default_rng(42)
    start = datetime(2023, 1, 1)
    rows = []
    for i in range(1096):
        current = start + timedelta(days=i)
        sales = 1800 + i * 1.1 + 290 * math.sin(2 * math.pi * current.weekday() / 7)
        sales += 100 * math.sin(2 * math.pi * i / 365.25) + rng.normal(0, 95)
        if current.day in (1, 15) and current.month in (1, 6, 11, 12):
            sales += 450
        rows.append([current.isoformat()[:10], round(sales, 2), "Online"])
    return ["date", "sales", "channel"], rows


def load_table(payload):
    if payload.get("sample") is True:
        headers, rows = sample_table()
        meta = {"filename": "synthetic_daily_sales.csv", "sheet": None, "sheets": [], "synthetic": True}
    elif "xlsx" in payload:
        try:
            raw = base64.b64decode(payload["xlsx"], validate=True)
            if len(raw) > MAX_FILE_BYTES:
                raise ForecastError("XLSX vượt giới hạn 2 MiB.", 413)
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                entries = archive.infolist()
                if len(entries) > 1000 or sum(item.file_size for item in entries) > 25 * 1024 * 1024:
                    raise ForecastError("XLSX giải nén quá lớn.", 413)
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            try:
                sheets = workbook.sheetnames
                if not sheets or len(sheets) > 40:
                    raise ForecastError("Workbook cần từ 1 đến 40 sheet.")
                selected = payload.get("sheet", sheets[0])
                if selected not in sheets:
                    raise ForecastError("Sheet đã chọn không tồn tại.")
                sheet = workbook[selected]
                if sheet.max_column and sheet.max_column > MAX_COLUMNS:
                    raise ForecastError("Tệp vượt giới hạn 60 cột.", 413)
                iterator = sheet.iter_rows(values_only=True)
                headers = next(iterator, [])
                rows = []
                for row in iterator:
                    rows.append(list(row))
                    if len(rows) > MAX_ROWS or len(rows) * len(headers) > MAX_CELLS:
                        raise ForecastError("Tệp vượt giới hạn 50.000 dòng / 1,5 triệu ô.", 413)
                meta = {"filename": token(payload.get("filename")) or "upload.xlsx", "sheet": selected,
                        "sheets": sheets, "synthetic": False}
            finally:
                workbook.close()
        except ForecastError:
            raise
        except Exception:
            raise ForecastError("Không đọc được XLSX. Hãy dùng workbook không mã hóa, có hàng tiêu đề.") from None
    else:
        text = payload.get("csv")
        if not isinstance(text, str) or not text.strip():
            raise ForecastError("Hãy tải CSV/XLSX hoặc chọn dữ liệu mẫu.", 400)
        if len(text.encode("utf-8")) > MAX_FILE_BYTES:
            raise ForecastError("CSV vượt giới hạn 2 MiB.", 413)
        if "\x00" in text:
            raise ForecastError("CSV chứa ký tự không hợp lệ.", 400)
        text = text.lstrip("\ufeff")
        try:
            delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
        except csv.Error:
            candidates = [(len(next(csv.reader(io.StringIO(text), delimiter=d), [])), d) for d in ",;\t|"]
            delimiter = max(candidates)[1]
        try:
            iterator = csv.reader(io.StringIO(text), delimiter=delimiter, strict=True)
            headers = next(iterator, [])
            rows = []
            for row in iterator:
                if len(row) > len(headers):
                    raise ForecastError("Số ô trong dòng lớn hơn số cột. Hãy kiểm tra dấu phân cách.", 400)
                rows.append(row + [None] * (len(headers) - len(row)))
                if len(rows) > MAX_ROWS or len(rows) * len(headers) > MAX_CELLS:
                    raise ForecastError("Tệp vượt giới hạn 50.000 dòng / 1,5 triệu ô.", 413)
        except csv.Error:
            raise ForecastError("CSV không hợp lệ; hãy kiểm tra dấu ngoặc kép.", 400) from None
        meta = {"filename": token(payload.get("filename")) or "upload.csv", "sheet": None,
                "sheets": [], "delimiter": delimiter, "synthetic": False}
    if not 2 <= len(headers) <= MAX_COLUMNS:
        raise ForecastError("Tệp cần từ 2 đến 60 cột và một hàng tiêu đề.", 400)
    original_rows = len(rows)
    rows = [row for row in rows if any(token(v).lower() not in MISSING for v in row)]
    if not rows:
        raise ForecastError("Tệp không có dòng dữ liệu.")
    kept = [i for i in range(len(headers)) if any(token(row[i]).lower() not in MISSING for row in rows)]
    used, names = set(), []
    renamed = 0
    for i in kept:
        name = token(headers[i]) or f"column_{i + 1}"
        candidate, suffix = name, 2
        while candidate in used:
            candidate = f"{name}_{suffix}"
            suffix += 1
        renamed += candidate != str(headers[i])
        used.add(candidate)
        names.append(candidate)
    cleaned, seen = [], set()
    for row in rows:
        values = [row[i] if isinstance(row[i], (date, datetime, int, float)) else token(row[i]) for i in kept]
        key = tuple(token(v) for v in values)
        if key not in seen:
            seen.add(key)
            cleaned.append(values)
    audit = {"input_rows": original_rows, "empty_rows": original_rows - len(rows),
             "empty_columns": len(headers) - len(kept), "renamed_headers": renamed,
             "duplicates_removed": len(rows) - len(cleaned)}
    return names, cleaned, meta, audit


def inspect(payload):
    names, rows, meta, audit = load_table(payload)
    columns = []
    for index, name in enumerate(names):
        values = [row[index] for row in rows[:1000] if token(row[index]).lower() not in MISSING]
        try:
            order, ambiguous = date_order(values)
        except ForecastError:
            order, ambiguous = "auto", True
        dates = sum(parse_date(v, order, excel="xlsx" in payload and isinstance(v, (date, datetime))) is not None for v in values)
        numbers = sum(numeric(v) is not None for v in values)
        dtype = "datetime" if values and dates / len(values) > .7 else "numeric" if values and numbers / len(values) > .7 else "text"
        date_hint = bool(re.search(r"date|datetime|timestamp|month|day|time|ngày", name, re.I))
        target_hint = bool(re.search(r"sales|revenue|quantity|amount|demand|units|turnover|doanh", name, re.I))
        columns.append({"name": name, "type": dtype, "date_score": dates / max(1, len(values)) + date_hint,
                        "target_score": numbers / max(1, len(values)) + target_hint, "ambiguous_dates": ambiguous})
    date_col = max(columns, key=lambda c: c["date_score"])
    targets = [c for c in columns if c != date_col]
    return {"success": True, "metadata": meta, "rows": len(rows), "columns": columns, "cleaning": audit,
            "preview": [[v.isoformat() if isinstance(v, (date, datetime)) else v for v in row] for row in rows[:10]],
            "suggested_date": date_col["name"] if date_col["date_score"] > .5 else None,
            "suggested_target": max(targets, key=lambda c: c["target_score"])["name"] if targets else None}


def infer_frequency(dates):
    deltas = [(b - a).total_seconds() for a, b in zip(dates, dates[1:])]
    if not deltas:
        return {"frequency": "daily", "interval": 1, "confidence": 0, "irregular": True}
    month_steps = [(b.year - a.year) * 12 + b.month - a.month for a, b in zip(dates, dates[1:])]
    aligned = len({(d.day, d.hour, d.minute) for d in dates}) == 1 or all(d.day == calendar.monthrange(d.year, d.month)[1] for d in dates)
    if aligned and all(step > 0 for step in month_steps):
        step, count = Counter(month_steps).most_common(1)[0]
        frequency = "quarterly" if step >= 3 and all(s % 3 == 0 for s in month_steps) else "monthly"
        confidence = count / len(month_steps)
    else:
        step, count = Counter(deltas).most_common(1)[0]
        if step < 86400:
            frequency = "hourly"
        elif step >= 604800 and all(d % 604800 == 0 for d in deltas):
            frequency = "weekly"
        else:
            frequency = "daily"
        confidence = count / len(deltas)
    # Normalize to unit calendar buckets. Larger source gaps become missing periods.
    return {"frequency": frequency, "interval": 1, "source_interval": float(step),
            "confidence": confidence, "irregular": confidence < .8}


def bucket(current, frequency):
    if frequency == "hourly":
        return current.replace(minute=0, second=0, microsecond=0)
    result = current.replace(hour=0, minute=0, second=0, microsecond=0)
    if frequency == "weekly":
        return result - timedelta(days=result.weekday())
    if frequency in ("monthly", "quarterly"):
        month = ((result.month - 1) // 3) * 3 + 1 if frequency == "quarterly" else result.month
        return result.replace(month=month, day=1)
    return result


def next_date(current, frequency):
    if frequency in ("monthly", "quarterly"):
        offset = 3 if frequency == "quarterly" else 1
        total = current.year * 12 + current.month - 1 + offset
        return current.replace(year=total // 12, month=total % 12 + 1, day=1)
    return current + {"hourly": timedelta(hours=1), "daily": timedelta(days=1), "weekly": timedelta(weeks=1)}[frequency]


def prepare(payload):
    names, rows, metadata, cleaning = load_table(payload)
    time_col, target_col = payload.get("date_column"), payload.get("target_column")
    if time_col not in names or target_col not in names or time_col == target_col:
        raise ForecastError("Hãy chọn hai cột thời gian và doanh số khác nhau.")
    group = payload.get("group_column")
    if group:
        if group not in names or group in (time_col, target_col) or "group_value" not in payload:
            raise ForecastError("Chọn cột nhóm và giá trị nhóm hợp lệ.")
        rows = [row for row in rows if token(row[names.index(group)]) == token(payload["group_value"])]
        if not rows:
            raise ForecastError("Nhóm đã chọn không có dữ liệu.")
    ti, yi = names.index(time_col), names.index(target_col)
    order, ambiguous = date_order([row[ti] for row in rows], payload.get("date_order", "auto"))
    if ambiguous:
        raise ForecastError("Ngày mơ hồ giữa DD/MM và MM/DD. Hãy chọn cách đọc ngày trước khi phân tích.")
    convention = payload.get("numeric_format", "auto")
    if convention not in ("auto", "decimal_dot", "decimal_comma"):
        raise ForecastError("Định dạng số không hợp lệ.")
    if convention == "auto" and any(re.fullmatch(r"[+-]?\d{1,3}[,.]\d{3}", token(row[yi])) for row in rows):
        raise ForecastError("Số như 1,234 / 1.234 có thể là thập phân hoặc hàng nghìn. Hãy chọn định dạng số.")
    points, bad_dates, bad_targets = [], 0, 0
    for row in rows:
        current = parse_date(row[ti], order, excel="xlsx" in payload)
        value = numeric(row[yi], convention)
        if current is None:
            bad_dates += 1
        elif value is None:
            bad_targets += 1
        else:
            points.append((current, value))
    if bad_dates > len(rows) / 2:
        raise ForecastError("Phần lớn giá trị thời gian không đọc được. Hãy kiểm tra cột và cách đọc ngày.")
    if bad_targets > len(rows) / 2 or len(points) < 3:
        raise ForecastError("Cần ít nhất 3 quan sát có ngày và doanh số hợp lệ.")
    cleaning.update(missing_dates=bad_dates, missing_targets=bad_targets, group_rows=len(rows))
    points.sort(key=lambda p: p[0])
    unique_dates = sorted({d for d, _ in points})
    inferred = infer_frequency(unique_dates)
    frequency = payload.get("frequency", "auto")
    if frequency == "auto":
        frequency = inferred["frequency"]
    if frequency not in PERIOD:
        raise ForecastError("Tần suất không hợp lệ.")
    sums = defaultdict(float)
    for current, value in points:
        sums[bucket(current, frequency)] += value
    if len(sums) < 3:
        raise ForecastError("Sau tổng hợp còn dưới 3 kỳ. Hãy chọn tần suất nhỏ hơn.")
    dates, values, missing = [], [], 0
    current, end = min(sums), max(sums)
    while current <= end:
        dates.append(current)
        values.append(sums.get(current, 0.0))
        missing += current not in sums
        if len(dates) > MAX_ROWS:
            raise ForecastError("Lưới thời gian vượt 50.000 kỳ. Hãy chọn tuần hoặc tháng.", 413)
        current = next_date(current, frequency)
    y = np.asarray(values, dtype=float)
    slope, intercept = np.polyfit(np.arange(len(y)), y, 1)
    scale = max(float(np.mean(np.abs(y))), 1e-9)
    trend_change = float(slope * (len(y) - 1))
    direction = "Increasing" if trend_change > .05 * scale else "Decreasing" if trend_change < -.05 * scale else "Stable"
    detrended = y - (intercept + slope * np.arange(len(y)))
    candidates = {"hourly": [24, 168], "daily": [7, 14, 30], "weekly": [4, 13, 52], "monthly": [3, 6, 12], "quarterly": [4]}[frequency]
    correlations = []
    for lag in candidates:
        if len(y) >= 3 * lag and np.std(detrended[:-lag]) > 1e-9 and np.std(detrended[lag:]) > 1e-9:
            correlations.append({"lag": lag, "correlation": float(np.corrcoef(detrended[:-lag], detrended[lag:])[0, 1])})
    seasonality = {"status": "Not enough data" if not correlations else "Detected" if max(c["correlation"] for c in correlations) >= .5 else "Weak", "correlations": correlations}
    median = float(np.median(detrended))
    mad = float(np.median(np.abs(detrended - median)))
    spikes = int(np.sum(np.abs(detrended - median) > max(3.5 * 1.4826 * mad, 1e-9))) if mad > 0 else 0
    warnings = []
    if metadata["synthetic"]:
        warnings.append("Dữ liệu mẫu mô phỏng; kết quả mô hình được tính thật từ dữ liệu này.")
    if missing:
        warnings.append(f"{missing} khoảng thời gian không có bản ghi được xem là doanh số 0.")
    if inferred["irregular"]:
        warnings.append("Tần suất nguồn không đều; hãy kiểm tra hoặc chọn lại tần suất tổng hợp.")
    if len(points) > len(unique_dates):
        warnings.append("Nhiều bản ghi cùng thời điểm đã được cộng doanh số.")
    if len(y) < 60:
        warnings.append("Lịch sử ngắn: dưới 12 kỳ chỉ dự báo thăm dò; 12–23 kỳ thăm dò; 24–59 kỳ backtesting hạn chế.")
    if np.std(y) / scale > 1:
        warnings.append("Doanh số biến động mạnh so với quy mô trung bình.")
    if np.mean(y == 0) > .2:
        warnings.append("Hơn 20% kỳ có doanh số 0; MAPE có thể không phù hợp.")
    if spikes:
        warnings.append(f"{spikes} kỳ có biến động bất thường; vẫn giữ trong dữ liệu huấn luyện.")
    if seasonality["status"] != "Detected":
        warnings.append("Chưa có bằng chứng mùa vụ mạnh qua tương quan của chuỗi đã loại xu hướng.")
    quality = {"status": "Insufficient" if len(y) < 12 else "Warning" if warnings else "Good", "observations": len(y),
               "start": dates[0].isoformat(), "end": dates[-1].isoformat(), "missing_periods": missing,
               "duplicates_removed": cleaning["duplicates_removed"], "aggregated_rows": len(points) - len(sums),
               "mean": float(np.mean(y)), "median": float(np.median(y)), "std": float(np.std(y)),
               "min": float(np.min(y)), "max": float(np.max(y)), "spikes": spikes,
               "trend": {"direction": direction, "change_per_period": float(slope),
                         "overall_percent": 100 * trend_change / abs(float(intercept)) if abs(intercept) > 1e-9 else None}}
    return {"dates": dates, "y": y, "metadata": {**metadata, "date_column": time_col, "target_column": target_col,
             "group_column": group, "group_value": payload.get("group_value"), "date_order": order, "numeric_format": convention},
            "cleaning": cleaning, "quality": quality, "frequency": {**inferred, "frequency": frequency,
            "inferred_frequency": inferred["frequency"], "overridden": payload.get("frequency", "auto") != "auto"},
            "seasonality": seasonality, "warnings": warnings}


def feature_schema(frequency, train_size):
    seasonal = {"hourly": [24, 48], "daily": [7, 14, 30], "weekly": [4, 13, 52], "monthly": [12, 24], "quarterly": [4]}[frequency]
    lags = [lag for lag in sorted(set([1, 2, 3] + seasonal)) if lag <= max(3, train_size // 3)]
    windows = [window for window in [3, 7, 12] if window <= max(3, train_size // 3)]
    names = ["time_index", "year", "month", "quarter", "day", "day_of_week", "week_of_year", "is_weekend", "hour"]
    names += [f"lag_{lag}" for lag in lags] + [f"rolling_mean_{w}" for w in windows] + [f"rolling_std_{w}" for w in windows if w in (3, 7)]
    return {"lags": lags, "windows": windows, "names": names, "warmup": max(lags + windows)}


def feature_at(history, current, index, schema):
    # history is an exclusive prefix y[:t]. Neither y[t] nor later values enter.
    if len(history) < schema["warmup"]:
        raise ForecastError("Chưa đủ lịch sử để tạo lag features.")
    result = [index, current.year, current.month, (current.month - 1) // 3 + 1, current.day,
              current.weekday(), current.isocalendar().week, int(current.weekday() >= 5), current.hour]
    result += [history[-lag] for lag in schema["lags"]]
    result += [float(np.mean(history[-w:])) for w in schema["windows"]]
    result += [float(np.std(history[-w:])) for w in schema["windows"] if w in (3, 7)]
    return result


def make_estimator(name):
    if name == "Linear Regression":
        return make_pipeline(StandardScaler(), LinearRegression())
    if name == "Ridge Regression":
        return make_pipeline(StandardScaler(), Ridge(alpha=10))
    if name == "Random Forest":
        return RandomForestRegressor(n_estimators=32, max_depth=9, min_samples_leaf=3, n_jobs=1, random_state=42)
    if name == "Gradient Boosting":
        return GradientBoostingRegressor(n_estimators=45, max_depth=2, min_samples_leaf=3, random_state=42)
    if name == "Hist Gradient Boosting" and HistGradientBoostingRegressor:
        # Disable shuffled internal early-stopping validation for temporal data.
        return HistGradientBoostingRegressor(max_iter=50, max_leaf_nodes=15, early_stopping=False, random_state=42)
    raise ForecastError("Mô hình không khả dụng.")


def fit_model(name, dates, y, schema):
    if name in MODELS[:2]:
        return None
    # Vectorized exclusive-prefix rolling sums use all rows with bounded cost.
    start = schema["warmup"]
    index = np.arange(start, len(y))
    current = dates[start:]
    columns = [index, [d.year for d in current], [d.month for d in current],
               [(d.month - 1) // 3 + 1 for d in current], [d.day for d in current],
               [d.weekday() for d in current], [d.isocalendar().week for d in current],
               [int(d.weekday() >= 5) for d in current], [d.hour for d in current]]
    columns += [y[index - lag] for lag in schema["lags"]]
    cumulative = np.concatenate(([0.0], np.cumsum(y)))
    # Center squares to avoid cancellation for large revenue scales.
    center = float(y[0])
    centered = y - center
    centered_sum = np.concatenate(([0.0], np.cumsum(centered)))
    squared_sum = np.concatenate(([0.0], np.cumsum(centered ** 2)))
    columns += [(cumulative[index] - cumulative[index - w]) / w for w in schema["windows"]]
    columns += [np.sqrt(np.maximum(0, (squared_sum[index] - squared_sum[index - w]) / w -
                                      ((centered_sum[index] - centered_sum[index - w]) / w) ** 2))
                for w in schema["windows"] if w in (3, 7)]
    X = np.column_stack(columns)
    if len(X) < 8:
        raise ForecastError("Chưa đủ quan sát huấn luyện sau tạo lag.")
    model = make_estimator(name)
    model.fit(X, y[schema["warmup"]:])
    return model


def recursive_predict(name, model, history, future_dates, schema, period, deadline=None):
    past = list(map(float, history))
    predictions = []
    for current in future_dates:
        if deadline and time.perf_counter() > deadline:
            raise TimeoutError("soft budget")
        if name == "Naive Forecast":
            value = past[-1]
        elif name == "Seasonal Naive":
            value = past[-period]
        else:
            value = float(model.predict(np.asarray([feature_at(past, current, len(past), schema)]))[0])
        if not math.isfinite(value) or abs(value) > 1e15:
            raise ForecastError("Dự báo mất ổn định.")
        predictions.append(value)
        past.append(value)
    return np.asarray(predictions)


def metrics(actual, prediction):
    actual, prediction = np.asarray(actual, dtype=float), np.asarray(prediction, dtype=float)
    error = actual - prediction
    denominator = np.abs(actual) + np.abs(prediction)
    smape = np.divide(2 * np.abs(error), denominator, out=np.zeros_like(error), where=denominator > 1e-9)
    total = float(np.sum((actual - np.mean(actual)) ** 2))
    return {"mae": float(np.mean(np.abs(error))), "rmse": float(np.sqrt(np.mean(error ** 2))),
            "smape": float(np.mean(smape) * 100),
            "mape": float(np.mean(np.abs(error / actual)) * 100) if np.all(np.abs(actual) > 1e-9) else None,
            "r2": 1 - float(np.sum(error ** 2)) / total if len(actual) >= 2 and total > 1e-9 else None}


def split_plan(size, horizon):
    if size < 12:
        return {"mode": "exploratory_no_validation", "holdout_start": size, "folds": [], "ranking_source": "none"}
    holdout_size = min(max(2, size // 5), 90)
    available = size - holdout_size
    count = 5 if size >= 300 else 3 if size >= 60 else 2 if size >= 24 else 1
    block = min(horizon, max(1, available // (count + 3)), 60)
    start = available - count * block
    folds = [{"train_end": start + i * block, "validation_start": start + i * block,
              "validation_end": start + (i + 1) * block} for i in range(count)]
    return {"mode": "normal" if size >= 60 else "limited_backtesting" if size >= 24 else "exploratory",
            "holdout_start": available, "folds": folds,
            "ranking_source": "backtest_rmse" if count > 1 else "inner_validation_rmse"}


def rank_models(records):
    return sorted([r for r in records if r["status"] == "SUCCESS" and r.get("backtest")],
                  key=lambda r: (r["backtest"]["rmse"], r["backtest"]["mae"]))


def baseline_warning(improvement):
    return "Mô hình được chọn chưa vượt Naive đáng kể (ngưỡng 5% RMSE); chưa đủ cơ sở gọi dự báo tốt." if improvement is None or improvement < 5 else None


def make_intervals(prediction, residuals, validated_steps):
    absolute = np.abs(np.asarray(residuals, dtype=float))
    if not len(absolute):
        return None
    q80, q95 = np.quantile(absolute, [.8, .95])
    output = []
    for i, value in enumerate(prediction):
        growth = math.sqrt(max(1, (i + 1) / max(1, validated_steps)))
        output.append({"lower_80": float(value - q80 * growth), "upper_80": float(value + q80 * growth),
                       "lower_95": float(value - q95 * growth), "upper_95": float(value + q95 * growth)})
    return output


def analyze_result(data):
    return {k: data[k] for k in ("metadata", "cleaning", "quality", "frequency", "seasonality", "warnings")} | {
        "success": True, "series": [{"date": d.isoformat(), "sales": float(v)} for d, v in zip(data["dates"], data["y"])],
        "raw_file_persisted": False}


def forecast(payload, soft_budget=45):
    started = time.perf_counter()
    horizon = payload.get("horizon", 30)
    if isinstance(horizon, bool) or not isinstance(horizon, int) or not 1 <= horizon <= 90:
        raise ForecastError("Số kỳ dự báo phải là số nguyên từ 1 đến 90.")
    data = prepare(payload)
    dates, y = data["dates"], data["y"]
    frequency = data["frequency"]["frequency"]
    period = PERIOD[frequency]
    plan = split_plan(len(y), horizon)
    folds = plan["folds"]
    schema = feature_schema(frequency, folds[0]["train_end"] if folds else len(y))
    deadline = started + soft_budget
    records, residuals_by_model = [], {}
    with threadpool_limits(limits=1):
        for name in MODELS:
            record = {"name": name, "status": "SKIPPED", "backtest": None, "holdout": None,
                      "features": schema["names"] if name not in MODELS[:2] else [], "lags": schema["lags"] if name not in MODELS[:2] else [],
                      "rolling_windows": schema["windows"] if name not in MODELS[:2] else [], "fit_time_ms": 0}
            records.append(record)
            smallest = folds[0]["train_end"] if folds else len(y)
            if name == "Seasonal Naive" and smallest < 2 * period:
                record.update(status="SKIPPED_INSUFFICIENT_HISTORY", notes=f"Cần ít nhất {2 * period} kỳ trong train.")
                continue
            if not folds and name != "Naive Forecast":
                record.update(status="SKIPPED_INSUFFICIENT_HISTORY", notes="Dưới 12 kỳ: chỉ dùng Naive thăm dò.")
                continue
            if name not in MODELS[:2] and smallest - schema["warmup"] < 8:
                record.update(status="SKIPPED_INSUFFICIENT_HISTORY", notes="Cần 8 dòng train sau lag warmup.")
                continue
            if name == "Hist Gradient Boosting" and HistGradientBoostingRegressor is None:
                record.update(status="SKIPPED", notes="sklearn không hỗ trợ estimator này.")
                continue
            if name != "Naive Forecast" and time.perf_counter() > deadline - 10:
                record.update(status="SKIPPED_TIME_BUDGET", notes="Giữ thời gian cho holdout và dự báo cuối.")
                continue
            try:
                scores, residuals, fit_time = [], [], 0
                for fold in folds:
                    end, stop = fold["train_end"], fold["validation_end"]
                    tick = time.perf_counter()
                    model = fit_model(name, dates[:end], y[:end], schema)
                    fit_time += time.perf_counter() - tick
                    pred = recursive_predict(name, model, y[:end], dates[end:stop], schema, period,
                                             deadline - 8 if name != "Naive Forecast" else None)
                    scores.append(metrics(y[end:stop], pred))
                    residuals.extend((y[end:stop] - pred).tolist())
                backtest = {key: float(np.mean([s[key] for s in scores])) if all(s[key] is not None for s in scores) else None
                            for key in ("mae", "rmse", "smape", "mape", "r2")} if scores else None
                if backtest:
                    backtest["rmse_std"] = float(np.std([s["rmse"] for s in scores]))
                    backtest["folds"] = scores
                record.update(status="SUCCESS", backtest=backtest, fit_time_ms=fit_time * 1000,
                              notes="Recursive block prediction; chỉ lịch sử train và dự báo trước đó.")
                residuals_by_model[name] = residuals
            except TimeoutError:
                record.update(status="SKIPPED_TIME_BUDGET", notes="Đã dừng model giữa các bước dự báo.")
            except ForecastError:
                record.update(status="FAILED_SAFE", notes="Mô hình thiếu lịch sử hoặc dự báo mất ổn định.")
            except Exception:
                logger.exception("sales_forecast model failed: %s", name)
                record.update(status="FAILED_SAFE", notes="Mô hình không hoàn tất; các mô hình khác tiếp tục.")
        ranked = rank_models(records)
        best = ranked[0] if ranked else records[0]
        for rank, record in enumerate(ranked, 1):
            record["rank"] = rank
        naive = records[0].get("backtest")
        for record in ranked:
            record["improvement_vs_naive"] = 100 * (naive["rmse"] - record["backtest"]["rmse"]) / naive["rmse"] if naive and naive["rmse"] > 1e-9 else None
        warning = baseline_warning(best.get("improvement_vs_naive"))
        if warning:
            data["warnings"].append(warning)
        # Selection is frozen before any holdout target is evaluated.
        holdout = {"dates": [], "actual": [], "predicted": [], "metrics": None, "used_for_ranking": False}
        cutoff = plan["holdout_start"]
        for record in [best] + [r for r in ranked if r is not best]:
            if cutoff >= len(y):
                break
            if record is not best and time.perf_counter() > deadline - 6:
                record["holdout_note"] = "Bỏ đánh giá phụ để giữ thời gian dự báo cuối."
                continue
            try:
                tick = time.perf_counter()
                model = fit_model(record["name"], dates[:cutoff], y[:cutoff], schema)
                record["fit_time_ms"] += (time.perf_counter() - tick) * 1000
                pred = recursive_predict(record["name"], model, y[:cutoff], dates[cutoff:], schema, period,
                                         deadline - 4 if record is not best else None)
                record["holdout"] = metrics(y[cutoff:], pred)
                if record is best:
                    residual = y[cutoff:] - pred
                    holdout.update(dates=[d.isoformat() for d in dates[cutoff:]], actual=y[cutoff:].tolist(),
                                   predicted=pred.tolist(), metrics=record["holdout"],
                                   residuals=residual.tolist(), mean_residual=float(np.mean(residual)),
                                   bias="underforecast" if np.mean(residual) > 0 else "overforecast" if np.mean(residual) < 0 else "neutral")
            except Exception:
                record["holdout_note"] = "Đánh giá holdout không hoàn tất; thứ hạng backtest vẫn giữ nguyên."
                if record is best:
                    data["warnings"].append("Holdout của model tốt nhất không hoàn tất.")
        future_dates, current = [], dates[-1]
        for _ in range(horizon):
            current = next_date(current, frequency)
            future_dates.append(current)
        try:
            tick = time.perf_counter()
            final_model = fit_model(best["name"], dates, y, schema)
            best["fit_time_ms"] += (time.perf_counter() - tick) * 1000
            prediction = recursive_predict(best["name"], final_model, y, future_dates, schema, period)
            forecast_model = best["name"]
        except Exception:
            final_model = None
            forecast_model = "Naive Forecast"
            prediction = recursive_predict(forecast_model, None, y, future_dates, schema, period)
            data["warnings"].append("Refit model đã chọn không hoàn tất; dự báo cuối dùng Naive an toàn.")
        if payload.get("clamp_nonnegative") is True and np.min(y) >= 0 and np.any(prediction < 0):
            prediction = np.maximum(prediction, 0)
            data["warnings"].append("Negative predictions were clipped to zero because historical sales are non-negative.")
        residuals = residuals_by_model.get(forecast_model, [])
        validated_steps = folds[0]["validation_end"] - folds[0]["validation_start"] if folds else 1
        bands = make_intervals(prediction, residuals, validated_steps)
        if bands is None:
            data["warnings"].append("Chưa có residual validation độc lập: không thể ước lượng khoảng dự báo.")
        elif len(residuals) < 30:
            data["warnings"].append("Dưới 30 residual backtest: khoảng 80%/95% chỉ là ước lượng rất hạn chế.")
        importance = []
        weights = getattr(final_model, "feature_importances_", None)
        importance_method = "tree impurity importance"
        if weights is None and final_model is not None and hasattr(final_model, "steps"):
            weights = np.abs(final_model.steps[-1][1].coef_)
            importance_method = "absolute coefficient on standardized features"
        if weights is not None:
            importance = sorted([{"feature": n, "value": float(w)} for n, w in zip(schema["names"], weights)], key=lambda item: -item["value"])[:10]
    rows = [{"date": d.isoformat(), "forecast": float(v), **(bands[i] if bands else {k: None for k in ("lower_80", "upper_80", "lower_95", "upper_95")})}
            for i, (d, v) in enumerate(zip(future_dates, prediction))]
    if horizon > len(y) / 3:
        data["warnings"].append("Horizon dài so với lịch sử; sai số đệ quy có thể tích lũy mạnh.")
    relative = best["backtest"]["rmse"] / max(float(np.mean(np.abs(y[:cutoff]))), 1e-9) if best.get("backtest") else math.inf
    stability = best["backtest"]["rmse_std"] / max(best["backtest"]["rmse"], 1e-9) if best.get("backtest") else math.inf
    reliability = "High" if len(y) >= 120 and relative < .15 and stability < .3 and horizon / len(y) <= .15 else "Medium" if len(y) >= 60 and relative < .4 and stability < .7 and horizon / len(y) <= .3 else "Low"
    previous = float(np.sum(y[-horizon:])) if horizon <= len(y) else None
    result = analyze_result(data)
    result.update(evaluation={**plan, "holdout_start_date": dates[cutoff].isoformat() if cutoff < len(y) else None,
                             "holdout_used_for_ranking": False, "shuffle": False,
                             "folds": [{**fold, "train_last_date": dates[fold["train_end"] - 1].isoformat(),
                                        "validation_first_date": dates[fold["validation_start"]].isoformat(),
                                        "validation_last_date": dates[fold["validation_end"] - 1].isoformat()} for fold in folds]},
                  models=ranked + [r for r in records if r not in ranked], best_model=best["name"], forecast_model=forecast_model,
                  holdout=holdout, forecast=rows, intervals={"label": "Estimated forecast interval", "method": "symmetric absolute backtest residual quantiles; sqrt widening beyond validated block",
                  "residual_count": len(residuals), "validated_steps": validated_steps, "guaranteed_coverage": False},
                  feature_importance=importance, importance_method=importance_method,
                  insights={"total": float(np.sum(prediction)), "average": float(np.mean(prediction)),
                            "max_period": rows[int(np.argmax(prediction))], "min_period": rows[int(np.argmin(prediction))],
                            "previous_total": previous, "change_percent": 100 * (float(np.sum(prediction)) - previous) / abs(previous) if previous is not None and abs(previous) > 1e-9 else None,
                            "reliability": reliability, "relative_backtest_rmse": relative if math.isfinite(relative) else None,
                            "stability_ratio": stability if math.isfinite(stability) else None},
                  refit_observations=len(y), duration_ms=(time.perf_counter() - started) * 1000,
                  partial=any(r["status"] in ("FAILED_SAFE", "SKIPPED_TIME_BUDGET") for r in records))
    return result


def dispatch(payload, isolated=True):
    if not isinstance(payload, dict):
        raise ForecastError("Request JSON cần là object.", 400)
    filename = payload.get("filename")
    if filename and (not isinstance(filename, str) or Path(filename).suffix.lower() not in (".csv", ".xlsx")):
        raise ForecastError("Chỉ hỗ trợ CSV và XLSX.", 400)
    action = payload.get("action")
    if action == "inspect":
        return inspect(payload)
    if action == "groups":
        names, rows, _, _ = load_table(payload)
        group = payload.get("group_column")
        if group not in names:
            raise ForecastError("Cột nhóm không hợp lệ.")
        values = sorted({token(row[names.index(group)]) for row in rows})
        if len(values) > 500:
            raise ForecastError("Cột có hơn 500 nhóm. Hãy lọc tệp trước khi tải.")
        return {"success": True, "values": values}
    if action == "analyze":
        return analyze_result(prepare(payload))
    if action != "forecast":
        raise ForecastError("Action không hợp lệ.", 400)
    if not isolated:
        return forecast(payload)
    env = os.environ.copy()
    env["PYTHONPATH"] = os.pathsep.join(dict.fromkeys([p for p in sys.path if isinstance(p, str) and p] + env.get("PYTHONPATH", "").split(os.pathsep)))
    try:
        completed = subprocess.run([sys.executable, "-B", str(Path(__file__).resolve()), "--worker"],
                                   input=json.dumps(payload, ensure_ascii=False).encode("utf-8"), stdout=subprocess.PIPE,
                                   stderr=None, timeout=55, env=env,
                                   creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0)
        if completed.returncode:
            raise ForecastError("Worker dự báo không hoàn tất; hãy thử lại với tệp nhỏ hơn.", 500)
        result = json.loads(completed.stdout)
        if "error" in result:
            raise ForecastError(result["error"], result.get("status", 500))
        return result
    except subprocess.TimeoutExpired:
        raise ForecastError("Đã dừng sau 55 giây. Hãy giảm dữ liệu hoặc horizon.", 504) from None


class handler(BaseHTTPRequestHandler):
    def log_message(self, *_args):
        pass

    def respond(self, status, body):
        raw = json.dumps(body, ensure_ascii=False, allow_nan=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("Content-Length", str(len(raw)))
        self.end_headers()
        self.wfile.write(raw)

    def do_GET(self):
        self.respond(200, {"available": True, "max_file_bytes": MAX_FILE_BYTES, "max_rows": MAX_ROWS, "raw_file_persisted": False})

    def do_POST(self):
        try:
            if self.headers.get("Content-Type", "").split(";", 1)[0].strip().lower() != "application/json":
                raise ForecastError("Request cần Content-Type application/json.", 400)
            try:
                length = int(self.headers.get("Content-Length", "0"))
            except ValueError:
                raise ForecastError("Content-Length không hợp lệ.", 400) from None
            if length > MAX_REQUEST_BYTES:
                raise ForecastError("Request vượt giới hạn 4 MiB.", 413)
            if length <= 0:
                raise ForecastError("Request rỗng.", 400)
            self.respond(200, dispatch(json.loads(self.rfile.read(length).decode("utf-8"))))
        except ForecastError as error:
            self.respond(error.status, {"success": False, "error": str(error)})
        except (UnicodeError, json.JSONDecodeError):
            self.respond(400, {"success": False, "error": "JSON/UTF-8 không hợp lệ."})
        except Exception:
            logger.exception("sales_forecast unexpected request error")
            self.respond(500, {"success": False, "error": "Dịch vụ dự báo gặp lỗi. Hãy thử lại với dữ liệu nhỏ hơn."})

    def do_PUT(self):
        self.respond(405, {"error": "Chỉ hỗ trợ GET và POST."})


if __name__ == "__main__":
    if "--worker" in sys.argv:
        try:
            output = forecast(json.loads(sys.stdin.buffer.read().decode("utf-8")))
        except ForecastError as error:
            output = {"error": str(error), "status": error.status}
        except Exception:
            logger.exception("sales_forecast unexpected worker error")
            output = {"error": "Không thể hoàn tất dự báo.", "status": 500}
        sys.stdout.buffer.write(json.dumps(output, ensure_ascii=False, allow_nan=False).encode("utf-8"))
    elif "--serve" in sys.argv:
        class Local(handler, SimpleHTTPRequestHandler):
            def do_GET(self):
                if self.path.split("?")[0] == "/api/sales-forecast":
                    return handler.do_GET(self)
                return SimpleHTTPRequestHandler.do_GET(self)

            def do_POST(self):
                if self.path != "/api/sales-forecast":
                    return self.respond(404, {"error": "Endpoint không tồn tại."})
                return handler.do_POST(self)
        os.chdir(Path(__file__).resolve().parents[1])
        port = int(os.environ.get("SALES_FORECAST_PORT", "8001"))
        print(f"Sales Forecasting: http://127.0.0.1:{port}/sales-forecasting.html", flush=True)
        ThreadingHTTPServer(("127.0.0.1", port), Local).serve_forever()
