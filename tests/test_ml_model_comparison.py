"""Real sklearn and HTTP/browser QA. Run python -B tests/test_ml_model_comparison.py.
Use --backend-only to omit browser. No external data or writes to old modules.
"""
import base64
import csv
import importlib.util
import io
import http.client
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import threading
import time
import unittest
import urllib.error
import urllib.request
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from unittest.mock import patch
import numpy as np
from sklearn.datasets import make_classification, make_regression
from openpyxl import Workbook, load_workbook

sys.dont_write_bytecode = True
ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('ml_compare', ROOT / 'api/ml-compare.py')
lab = importlib.util.module_from_spec(spec)
spec.loader.exec_module(lab)


def csv_text(headers, rows):
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    writer.writerows(rows)
    return output.getvalue()


def fixture(task='classification', count=120, classes=2, mixed=True):
    if task == 'classification':
        X, y = make_classification(n_samples=count, n_features=4, n_informative=3, n_redundant=0,
                                    n_classes=classes, n_clusters_per_class=1, random_state=47)
    else:
        X, y = make_regression(n_samples=count, n_features=4, noise=15, random_state=47)
    headers = ['x0','x1','x2','x3'] + (['segment','customer_id','event_date'] if mixed else []) + ['target']
    rows = []
    for i, (x, target) in enumerate(zip(X, y)):
        values = list(x)
        if mixed:
            if i % 13 == 0:
                values[0] = ''
            values += ['' if i % 17 == 0 else ['north','south','west'][i % 3], 'C' + str(i), f'2025-{i // 28 + 1:02d}-{i % 28 + 1:02d}']
        rows.append(values + [int(target) if task == 'classification' else target])
    return dict(action='benchmark',csv=csv_text(headers,rows),target='target',task=task,
                features=['x0','x1','x2','x3'] + (['segment'] if mixed else []),split='random')


def excel_fixture(task='classification', mixed=True, multiple=False, count=120):
    payload = fixture(task, mixed=mixed, count=count)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Dữ liệu'
    for row in csv.reader(io.StringIO(payload.pop('csv'))):
        sheet.append([None if value == '' else float(value) if lab.number(value) is not None else value for value in row])
    if multiple:
        other = workbook.create_sheet('Hồi quy')
        for row in csv.reader(io.StringIO(fixture('regression', mixed=False)['csv'])):
            other.append([float(value) if lab.number(value) is not None else value for value in row])
    buffer = io.BytesIO(); workbook.save(buffer); workbook.close()
    payload['xlsx'] = base64.b64encode(buffer.getvalue()).decode('ascii')
    return payload


def weak_fixture():
    payload=fixture('regression',count=60,mixed=False)
    rows=list(csv.reader(io.StringIO(payload['csv'])))
    rng=np.random.default_rng(20)
    for row,value in zip(rows[1:],rng.normal(size=60)):row[-1]=str(value)
    payload['csv']=csv_text(rows[0],rows[1:])
    return payload


class Backend(unittest.TestCase):
    def test_worker_environment_preserves_order_existing_paths_and_separator(self):
        with tempfile.TemporaryDirectory() as directory:
            existing = str(Path(directory) / 'existing')
            with patch.object(lab.sys, 'path', [directory, '', directory, None]), \
                 patch.dict(lab.os.environ, {'PYTHONPATH': os.pathsep.join([existing, directory, '', existing])}):
                before = os.environ.copy()
                env = lab._worker_environment()
                self.assertEqual(env['PYTHONPATH'], os.pathsep.join([directory, existing]))
                self.assertEqual(env['PYTHONPATH'].split(os.pathsep), [directory, existing])
                self.assertEqual(os.environ.copy(), before)
                self.assertEqual({k: v for k, v in env.items() if k != 'PYTHONPATH'},
                                 {k: v for k, v in before.items() if k != 'PYTHONPATH'})

    def test_child_imports_module_from_injected_parent_path(self):
        with tempfile.TemporaryDirectory() as directory:
            Path(directory, 'worker_runtime_path_probe.py').write_text('VALUE = 47\n', encoding='utf-8')
            with patch.object(lab.sys, 'path', [directory] + sys.path):
                result = subprocess.run([sys.executable, '-B', '-c',
                                         'import worker_runtime_path_probe as probe; print(probe.VALUE)'],
                                        env=lab._worker_environment(), capture_output=True, timeout=15,
                                        creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0)
            self.assertEqual(result.returncode, 0, result.stderr.decode())
            self.assertEqual(result.stdout.strip(), b'47')

    def test_worker_exception_logged_and_client_sanitized(self):
        from types import SimpleNamespace
        request = {'action': 'benchmark', 'csv': 'RAW_REQUEST_MUST_NOT_BE_LOGGED'}
        output = io.BytesIO()
        with patch.object(lab.sys, 'stdin', SimpleNamespace(buffer=io.BytesIO(json.dumps(request).encode()))), \
             patch.object(lab.sys, 'stdout', SimpleNamespace(buffer=output)), \
             patch.object(lab, 'benchmark', side_effect=RuntimeError('PRIVATE_WORKER_FAILURE')), \
             self.assertLogs(lab.logger, level='ERROR') as logs:
            lab.run_worker()
        logged = '\n'.join(logs.output)
        self.assertIn('Traceback', logged)
        self.assertIn('RuntimeError: PRIVATE_WORKER_FAILURE', logged)
        self.assertNotIn('RAW_REQUEST_MUST_NOT_BE_LOGGED', logged)
        response = json.loads(output.getvalue())
        self.assertEqual(response['status'], 500)
        self.assertNotIn('PRIVATE_WORKER_FAILURE', str(response))
        self.assertNotIn('Traceback', str(response))

    def test_worker_startup_stderr_reaches_server_not_client(self):
        completed = subprocess.CompletedProcess([], 1, stdout=b'')
        with patch.object(lab.subprocess, 'run', return_value=completed) as run, \
             self.assertLogs(lab.logger, level='ERROR') as logs:
            with self.assertRaises(lab.LabError) as error:
                lab.dispatch(fixture())
        self.assertIsNone(run.call_args.kwargs['stderr'])
        self.assertEqual(run.call_args.kwargs['env'], lab._worker_environment())
        self.assertEqual(run.call_args.kwargs['timeout'], 55)
        self.assertEqual(error.exception.status, 500)
        self.assertIn('worker exited with code=1', '\n'.join(logs.output))
        self.assertNotIn('Traceback', str(error.exception))

    def test_xlsx_without_pandas_and_excel_empty_cells(self):
        import builtins
        workbook = Workbook(); sheet = workbook.active
        sheet.append([' x ', None, 'target'])
        sheet.append([1, '=1+1', 'yes'])  # Uncached formula is missing with data_only.
        sheet.append([2, '#DIV/0!', 'no'])
        sheet.append([None, None, None])
        sheet.append([3, 'NA', 'yes'])
        sheet.cell(10, 6).value = None  # Formatted trailing blanks are not columns.
        sheet.cell(10, 6).number_format = '0.00'
        buffer = io.BytesIO(); workbook.save(buffer); workbook.close()
        original = builtins.__import__
        def no_pandas(name, *args, **kwargs):
            if name.split('.')[0] == 'pandas':
                raise AssertionError('Production XLSX must not import pandas')
            return original(name, *args, **kwargs)
        with patch.object(builtins, '__import__', no_pandas):
            headers, rows, audit = lab.load_data({'xlsx': base64.b64encode(buffer.getvalue()).decode()})
        self.assertEqual(headers, ['x', 'column_2', 'target'])
        self.assertEqual(rows, [('1', None, 'yes'), ('2', None, 'no'), ('3', None, 'yes')])
        self.assertEqual(audit['blank_rows'], 1)
        self.assertEqual(audit['short_rows_padded'], 0)

    def test_xlsx_stale_dimensions_do_not_truncate_cells(self):
        import zipfile
        payload = excel_fixture()
        output = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(base64.b64decode(payload['xlsx']))) as source:
            with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as dest:
                for entry in source.infolist():
                    content = source.read(entry)
                    if entry.filename == 'xl/worksheets/sheet1.xml':
                        import re
                        content = re.sub(br'<dimension ref="[^"]+"', b'<dimension ref="A1:A1"', content)
                    dest.writestr(entry, content)
        payload['xlsx'] = base64.b64encode(output.getvalue()).decode()
        headers, rows, _ = lab.load_data(payload)
        self.assertEqual(len(headers), 8)
        self.assertEqual(len(rows), 120)

    def test_xlsx_numeric_mixed_missing_and_shared_pipeline(self):
        for task, mixed in [('classification',False),('classification',True),('regression',True)]:
            with self.subTest(task=task,mixed=mixed):
                payload = excel_fixture(task,mixed)
                headers,rows,audit = lab.load_data(payload)
                summary = lab.dispatch({**payload,'action':'inspect','target':'target'})
                self.assertEqual(audit['selected_sheet'],'Dữ liệu')
                self.assertEqual(summary['rows'],120)
                self.assertEqual(summary['columns'][-1]['suggested_task'],task)
                self.assertEqual(summary['missing'] > 0,mixed)
                result = lab.benchmark(payload)
                self.assertEqual(result['selection_source'],'train_cv')
                self.assertEqual(len(result['models']),7 if task=='classification' else 8)
                self.assertFalse(result['raw_file_persisted'])
                self.assertNotIn('xlsx',result)
                self.assertNotIn('preview',result['summary'])

    def test_xlsx_multiple_sheets_and_explicit_selection(self):
        payload = excel_fixture(multiple=True)
        summary = lab.dispatch({**payload,'action':'inspect'})
        self.assertEqual(summary['audit']['sheets'],['Dữ liệu','Hồi quy'])
        self.assertEqual(summary['columns'][-1]['suggested_task'],'classification')
        selected = {**payload,'sheet':'Hồi quy','task':'regression','features':['x0','x1','x2','x3']}
        self.assertEqual(lab.dispatch({**selected,'action':'inspect'})['columns'][-1]['suggested_task'],'regression')
        result = lab.benchmark(selected)
        self.assertEqual(result['summary']['audit']['selected_sheet'],'Hồi quy')
        with self.assertRaises(lab.LabError): lab.load_data({**payload,'sheet':'missing'})
        with self.assertRaises(lab.LabError): lab.prepare({**selected,'target':'absent'})

    def test_xlsx_cleaning_and_native_dates(self):
        from datetime import datetime
        payload=excel_fixture()
        workbook=load_workbook(io.BytesIO(base64.b64decode(payload['xlsx'])))
        sheet=workbook.active
        sheet.cell(2,7,datetime(2025,1,1))
        sheet.append([cell.value for cell in sheet[2]])
        sheet.append([None]*8)
        sheet.append([1,2,3,4,'north','C-missing','2025-12-31',None])
        buffer=io.BytesIO();workbook.save(buffer);workbook.close()
        payload['xlsx']=base64.b64encode(buffer.getvalue()).decode()
        summary=lab.dispatch({**payload,'action':'inspect'})
        self.assertEqual(summary['audit']['duplicates_removed'],1)
        self.assertEqual(summary['audit']['blank_rows'],1)
        self.assertTrue(summary['columns'][5]['id_like'])
        self.assertTrue(summary['columns'][6]['date'])
        self.assertEqual(lab.prepare(payload)['summary']['audit']['target_rows_removed'],1)

    def test_xlsx_invalid_and_resource_limits(self):
        for encoded in ['not base64',base64.b64encode(b'a,b\n1,2').decode(),base64.b64encode(b'PK\x03\x04broken').decode()]:
            with self.assertRaises(lab.LabError) as caught:lab.load_data({'xlsx':encoded})
            self.assertIn(caught.exception.status,(400,422))
            self.assertNotIn('Traceback',str(caught.exception))
        for rows,columns in [(20002,2),(2,101)]:
            workbook=Workbook();sheet=workbook.active
            sheet.cell(rows,columns,'too large')
            buffer=io.BytesIO();workbook.save(buffer);workbook.close()
            with self.assertRaises(lab.LabError) as caught:lab.load_data({'xlsx':base64.b64encode(buffer.getvalue()).decode()})
            self.assertEqual(caught.exception.status,413)
        # All sheets are checked, including sheets not selected by the user.
        workbook=Workbook();workbook.active.append(['x','target']);workbook.active.append([1,0])
        workbook.create_sheet('Oversized').cell(20002,1,'large')
        buffer=io.BytesIO();workbook.save(buffer);workbook.close()
        with self.assertRaises(lab.LabError) as caught:lab.load_data({'xlsx':base64.b64encode(buffer.getvalue()).decode(),'sheet':'Sheet'})
        self.assertEqual(caught.exception.status,413)
        with self.assertRaises(lab.LabError) as caught:lab.load_data({'xlsx':'A'*(4*((lab.MAX_CSV_BYTES+2)//3)+1)})
        self.assertEqual(caught.exception.status,413)

    def test_small_data_and_low_quality_do_not_change_cv_recommendation(self):
        self.assertTrue(any('khá nhỏ' in w for w in lab.prepare(fixture(count=60))['warnings']))
        self.assertTrue(any('rất nhỏ' in w for w in lab.prepare(fixture(count=35))['warnings']))
        original=lab.evaluate
        for task in ['classification','regression']:
            def weak_test(pipe,X,y,current_task):
                values,pred,elapsed,curves=original(pipe,X,y,current_task)
                if len(y)==30:values['f1' if current_task=='classification' else 'r2']=.01 if current_task=='classification' else -.5
                return values,pred,elapsed,curves
            with patch.object(lab,'evaluate',weak_test):result=lab.benchmark(fixture(task,mixed=False))
            self.assertEqual(result['quality']['level'],'low')
            self.assertEqual(result['recommended'],result['models'][0]['name'])
            self.assertEqual(result['selection_source'],'train_cv')
            self.assertFalse(result['methodology']['test_used_for_selection'])

        def near_zero(pipe,X,y,current_task):
            values,pred,elapsed,curves=original(pipe,X,y,current_task)
            if len(y)==30:values['r2']=.05
            return values,pred,elapsed,curves
        with patch.object(lab,'evaluate',near_zero):result=lab.benchmark(fixture('regression',mixed=False))
        self.assertEqual(result['quality']['level'],'caution')

    def test_binary_mixed_missing_and_metrics(self):
        payload = fixture()
        result = lab.benchmark(payload)
        self.assertEqual(len(result['models']),7)
        self.assertFalse(result['raw_file_persisted'])
        self.assertNotIn('preview',result['summary'])
        self.assertEqual(result['methodology']['cv'],'StratifiedKFold')
        self.assertEqual(result['methodology']['folds'],5)
        self.assertEqual(result['recommended'],result['models'][0]['name'])
        for model in result['models']:
            self.assertEqual(sum(map(sum,model['confusion']['matrix'])),30)
            self.assertEqual(len(model['cv_scores']),5)
            self.assertTrue(0 <= model['test']['f1'] <= 1)
            self.assertTrue(model['fit_time_ms'] > 0 and model['predict_time_ms'] > 0)
            self.assertTrue(model['model_size_bytes'] > 0)
            self.assertIn('classification_report',model)
            self.assertTrue(model['curves']['roc'] and model['curves']['pr'])
        json.dumps(lab.sanitize(result),allow_nan=False)

    def test_multiclass_numeric(self):
        result=lab.benchmark(fixture(classes=3,mixed=False))
        self.assertEqual(len(result['models']),7)
        for model in result['models']:
            self.assertEqual(len(model['confusion']['labels']),3)
            self.assertIsNone(model['curves'])

    def test_regression(self):
        result=lab.benchmark(fixture('regression'))
        self.assertEqual(len(result['models']),8)
        self.assertEqual(result['methodology']['cv'],'KFold')
        for model in result['models']:
            self.assertAlmostEqual(model['test']['rmse']**2,model['test']['mse'],places=6)
            actual,predicted=model['scatter'][0]
            self.assertAlmostEqual(model['residuals'][0][1],actual-predicted)

    def test_ranking_ignores_test(self):
        rows=[dict(name='A',cv_mean=.8,cv_std=.1,cv_fit_ms=10,test={'f1':.01}),
              dict(name='B',cv_mean=.7,cv_std=.01,cv_fit_ms=1,test={'f1':1})]
        self.assertEqual(lab.rank_models(rows,'classification')[0]['name'],'A')
        rows[0]['test']['f1']=1; rows[1]['test']['f1']=0
        self.assertEqual(lab.rank_models(rows,'classification')[0]['name'],'A')
        rows[0]['cv_mean']=10; rows[1]['cv_mean']=11
        rows[0]['test']={'rmse':999}; rows[1]['test']={'rmse':0}
        self.assertEqual(lab.rank_models(rows,'regression')[0]['name'],'A')

    def test_actual_benchmark_keeps_cv_choice_when_test_winner_changes(self):
        payload=fixture(mixed=False)
        # Full fixture ties NB and SVM on both CV mean and std. Its legitimate
        # timing tie-break can change between runs, unrelated to held-out scores.
        # Use two real estimators with distinct CV scores to isolate this invariant.
        catalogue=[m for m in lab.models('classification') if m[0] in ('Gaussian Naive Bayes','Logistic Regression')]
        with patch.object(lab,'models',return_value=catalogue):baseline=lab.benchmark(payload)
        self.assertNotEqual(baseline['models'][0]['cv_mean'],baseline['models'][1]['cv_mean'])
        winner=baseline['recommended']
        original=lab.evaluate
        def altered(pipe,X,y,task):
            values,pred,time_ms,curves=original(pipe,X,y,task)
            if len(y)==30:
                # Forced held-out ranking contradiction, only in this test.
                values['f1']=0 if type(pipe.named_steps['model']).__name__==baseline['models'][0]['algorithm_type'] else 1
            return values,pred,time_ms,curves
        with patch.object(lab,'evaluate',altered),patch.object(lab,'models',return_value=catalogue):
            result=lab.benchmark(payload)
        self.assertEqual(result['recommended'],winner)
        self.assertEqual(result['models'][0]['test']['f1'],0)
        self.assertEqual(max(r['test']['f1'] for r in result['models']),1)

    def test_preprocessing_fits_train_only_and_inside_folds(self):
        data=lab.prepare(fixture())
        X=data['X'].copy(); train,test=data['train'],data['test']
        X[test,0]=1000000; X[test,4]='TEST_ONLY_CATEGORY'
        p=lab.preprocessor(data,True).fit(X[train])
        numeric=p.named_transformers_['numeric']
        self.assertAlmostEqual(numeric.named_steps['imputer'].statistics_[0],np.nanmedian(np.asarray(X[train,0],dtype=float)))
        self.assertTrue(numeric.named_steps['scaler'].mean_[0] < 100)
        self.assertNotIn('TEST_ONLY_CATEGORY',p.named_transformers_['categorical'].named_steps['onehot'].categories_[0])
        self.assertEqual(p.transform(X[test]).shape[0],len(test))
        a,b=data['splits'][0]; X[train[b],4]='VALIDATION_ONLY'
        folded=lab.preprocessor(data,True).fit(X[train[a]])
        self.assertNotIn('VALIDATION_ONLY',folded.named_transformers_['categorical'].named_steps['onehot'].categories_[0])
        # Observe the real benchmark fit path, including every CV fit and final fit.
        original=lab.ColumnTransformer.fit_transform
        seen=[]
        def spy(transformer,array,*args,**kwargs):
            seen.append(len(array))
            return original(transformer,array,*args,**kwargs)
        with patch.object(lab.ColumnTransformer,'fit_transform',spy):
            lab.benchmark(fixture())
        expected=[len(a) for a,_ in data['splits']]+[len(train)]
        self.assertEqual(seen,expected*7)

    def test_csv_quality(self):
        text='\ufeff a ,a, target\n1,NA,yes\n1,NA,yes\n\n2,null,no\n3\n'
        headers,rows,audit=lab.parse_csv(text)
        self.assertEqual(headers,['a','a__2','target'])
        self.assertEqual(audit['duplicates_removed'],1)
        self.assertEqual(audit['blank_rows'],1)
        self.assertEqual(audit['short_rows_padded'],1)
        self.assertIsNone(rows[-1][-1])
        with self.assertRaises(lab.LabError):lab.parse_csv('a,b\n1,2,3')

    def test_ambiguous_dates_never_onehot(self):
        p=fixture();headers,rows,audit=lab.load_data(p)
        i=headers.index('event_date')
        rows=[tuple('01/02/2025' if j==i else v for j,v in enumerate(r)) for r in rows]
        p['csv']=csv_text(headers,rows);p['features']+=['event_date']
        self.assertTrue(lab.profile(headers,rows,audit)['columns'][i]['date'])
        self.assertNotIn('event_date',lab.prepare(p)['names'])
        with self.assertRaises(lab.LabError):lab.prepare({**p,'date_features':'extract'})
        self.assertEqual(lab.date('31/12/2025').month,12)
        self.assertEqual(lab.date('12/31/2025').month,12)

    def test_target_validation(self):
        p=fixture()
        for target in ('absent','customer_id','event_date'):
            with self.subTest(target=target),self.assertRaises(lab.LabError):lab.prepare({**p,'target':target})
        for value in ('','same'):
            p['csv']=csv_text(['x','target'],[[i,value] for i in range(60)])
            p['features']=['x']
            with self.assertRaises(lab.LabError):lab.prepare(p)
        with self.assertRaises(lab.LabError):lab.prepare({**fixture(),'features':['target']})
        with self.assertRaises(lab.LabError):lab.prepare({**fixture(),'task':'invalid'})

    def test_id_and_detection(self):
        headers,rows,audit=lab.load_data(fixture())
        summary=lab.profile(headers,rows,audit)
        by_name={c['name']:c for c in summary['columns']}
        self.assertTrue(by_name['customer_id']['id_like'])
        self.assertTrue(by_name['event_date']['date'])
        self.assertEqual(by_name['target']['suggested_task'],'classification')
        self.assertFalse(by_name['x1']['id_like'])
        headers,rows,audit=lab.load_data(fixture('regression'))
        self.assertEqual(lab.profile(headers,rows,audit)['columns'][-1]['suggested_task'],'regression')

    def test_temporal_and_date_features(self):
        p=fixture('regression')
        p.update(split='temporal',time_column='event_date',date_features='extract')
        p['features']+=['event_date']
        data=lab.prepare(p)
        self.assertEqual(data['train'].tolist(),list(range(90)))
        self.assertEqual(data['test'].tolist(),list(range(90,120)))
        for train,valid in data['splits']:self.assertLess(max(train),min(valid))
        self.assertIn('event_date__day_of_week',data['names'])
        result=lab.benchmark(p)
        self.assertEqual(result['methodology']['cv'],'TimeSeriesSplit')
        self.assertEqual(len(result['models']),8)
        p=fixture();p.update(split='temporal',time_column='event_date')
        self.assertEqual(lab.benchmark(p)['methodology']['cv'],'TimeSeriesSplit')
        with self.assertRaises(lab.LabError):lab.prepare({**p,'time_column':'customer_id'})

    def test_small_and_large_guards(self):
        result=lab.benchmark(fixture(count=35,mixed=False))
        self.assertEqual(len(result['models']),3)
        self.assertEqual(len(result['skipped']),4)
        with self.assertRaises(lab.LabError):lab.prepare(fixture(count=15,mixed=False))
        with self.assertRaises(lab.LabError) as error:lab.parse_csv('a,b\n'+'x'*(lab.MAX_CSV_BYTES+1))
        self.assertEqual(error.exception.status,413)
        with self.assertRaises(lab.LabError) as error:lab.parse_csv('a,b\n'+'1,2\n'*20001)
        self.assertEqual(error.exception.status,413)
        with self.assertRaises(lab.LabError):lab.parse_csv(','.join('x'+str(i) for i in range(101))+'\n'+','.join('1' for _ in range(101)))
        with patch.object(lab.subprocess,'run',side_effect=subprocess.TimeoutExpired('worker',55)):
            with self.assertRaises(lab.LabError) as error:lab.dispatch(fixture())
            self.assertEqual(error.exception.status,504)

    def test_leakage(self):
        p=fixture();headers,rows,audit=lab.load_data(p)
        headers+=['copy'];rows=[tuple(r)+(r[-1],) for r in rows]
        p['csv']=csv_text(headers,rows);p['features']+=['copy']
        self.assertTrue(lab.leakage_warnings(headers,rows,'target',['copy']))
        with self.assertRaises(lab.LabError):lab.prepare({**p,'ack_leakage':True})

    def test_high_cardinality_bounded(self):
        p=fixture();headers,rows,audit=lab.load_data(p)
        headers+=['category'];rows=[tuple(r)+('G'+str(i),) for i,r in enumerate(rows)]
        p['csv']=csv_text(headers,rows);p['features']+=['category']
        data=lab.prepare(p);transformed=lab.preprocessor(data,True).fit_transform(data['X'][data['train']])
        self.assertLessEqual(transformed.shape[1],4+3+32)
        self.assertTrue(any('OHE' in w for w in data['warnings']))

    def test_all_demos(self):
        for name in ['iris','wine','breast_cancer','digits','diabetes']:
            with self.subTest(name=name):
                p={'action':'benchmark','demo':name,'target':'target','task':'regression' if name=='diabetes' else 'classification'}
                headers,_,_=lab.load_data(p);p['features']=headers[:-1]
                result=lab.benchmark(p)
                self.assertEqual(len(result['models']),8 if name=='diabetes' else 7)
                print('PASS demo:',name,result['recommended'],flush=True)


class HTTP(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.server=ThreadingHTTPServer(('127.0.0.1',0),lab.handler)
        threading.Thread(target=cls.server.serve_forever,daemon=True).start()

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown();cls.server.server_close()

    def request(self,body,headers=None,method='POST'):
        conn=http.client.HTTPConnection('127.0.0.1',self.server.server_port,timeout=65)
        try:
            conn.request(method,'/api/ml-compare',body,headers or {'Content-Type':'application/json'})
            response=conn.getresponse()
            return response.status,json.loads(response.read()),dict(response.getheaders())
        finally:conn.close()

    def test_http_errors_and_privacy(self):
        status,body,headers=self.request(None,method='GET')
        self.assertEqual(status,200);self.assertEqual(headers['Cache-Control'],'no-store')
        self.assertFalse(body['raw_file_persisted'])
        for raw,status in [(b'{',400),(json.dumps({'action':'inspect','csv':'one\n1'}).encode(),400),
                           (json.dumps({**fixture(),'target':'absent'}).encode(),422)]:
            actual,body,_=self.request(raw)
            self.assertEqual(actual,status);self.assertIn('error',body);self.assertNotIn('Traceback',body['error'])
        status,_,_=self.request(None,{'Content-Type':'application/json','Content-Length':str(lab.MAX_REQUEST_BYTES+1)})
        self.assertEqual(status,413)
        status,_,_=self.request(b'x',{'Content-Type':'text/plain'})
        self.assertEqual(status,400)
        status,_,_=self.request(b'',method='PUT')
        self.assertEqual(status,405)

    def test_real_worker_training(self):
        status,result,_=self.request(json.dumps(fixture(mixed=False)).encode())
        self.assertEqual(status,200);self.assertEqual(len(result['models']),7)
        self.assertFalse(result['raw_file_persisted']);self.assertNotIn('csv',result)
        self.assertNotIn('preview',result['summary'])

    def test_deployment_http_fixtures_and_wall_clock(self):
        cases=[('small binary','classification',120,2),
               ('multiclass','classification',150,3),
               ('medium binary','classification',2000,2),
               ('regression','regression',500,2)]
        for name,task,count,classes in cases:
            with self.subTest(case=name):
                payload=fixture(task,count=count,classes=classes,mixed=False)
                started=time.perf_counter()
                status,result,headers=self.request(json.dumps(payload).encode())
                elapsed=time.perf_counter()-started
                self.assertEqual(status,200,result.get('error'))
                self.assertEqual(headers['Content-Type'],'application/json; charset=utf-8')
                self.assertEqual(result['selection_source'],'train_cv')
                self.assertFalse(result['methodology']['test_used_for_selection'])
                self.assertFalse(result['raw_file_persisted'])
                self.assertEqual(len(result['models']),7 if task=='classification' else 8)
                print(json.dumps({'http_fixture':name,'rows':count,'status':status,
                                  'wall_clock_seconds':round(elapsed,3),'benchmark_ms':round(result['duration_ms'],1),
                                  'models':len(result['models']),'skipped':result['skipped']}),flush=True)

    def test_csv_limit_and_sanitized_500_over_http(self):
        payload={'action':'inspect','csv':'a,b\n'+'x'*(lab.MAX_CSV_BYTES+1)}
        status,result,_=self.request(json.dumps(payload).encode())
        self.assertEqual(status,413)
        self.assertIn('2 MiB',result['error'])
        # Deliberate internal failure: error details must not reach the client.
        with patch.object(lab,'dispatch',side_effect=RuntimeError('PRIVATE_INTERNAL_TRACE')), self.assertLogs(lab.logger, level='ERROR') as logs:
            status,result,_=self.request(json.dumps({'action':'inspect','csv':'a,b\n1,2'}).encode())
        self.assertIn('RuntimeError: PRIVATE_INTERNAL_TRACE', '\n'.join(logs.output))
        self.assertIn('Traceback', '\n'.join(logs.output))
        self.assertEqual(status,500)
        self.assertNotIn('PRIVATE_INTERNAL_TRACE',json.dumps(result))
        self.assertNotIn('Traceback',json.dumps(result))

    def test_content_type_contract(self):
        body=json.dumps({'action':'inspect','csv':'a,b\n1,2'}).encode()
        for content_type in ('application/json','Application/JSON','APPLICATION/JSON','APPLICATION/JSON; charset=UTF-8','application/json; charset=utf-8','Application/JSON; charset=UTF-8'):
            with self.subTest(content_type=content_type):
                status,_,_=self.request(body,{'Content-Type':content_type})
                self.assertEqual(status,200)
        status,_,_=self.request(body,{'Content-Type':'application/jsonp'})
        self.assertEqual(status,400)

    def test_excel_http_and_sanitized_invalid_upload(self):
        payload=excel_fixture('regression',multiple=True)
        payload.update(sheet='Hồi quy',features=['x0','x1','x2','x3'])
        status,result,_=self.request(json.dumps(payload).encode())
        self.assertEqual(status,200,result.get('error'))
        self.assertEqual(result['summary']['audit']['selected_sheet'],'Hồi quy')
        self.assertEqual(len(result['models']),8)
        status,result,_=self.request(json.dumps({'action':'inspect','xlsx':'invalid'}).encode())
        self.assertEqual(status,422)
        self.assertNotIn('Traceback',result['error'])

    def test_deployment_source_contract(self):
        from html.parser import HTMLParser
        class References(HTMLParser):
            def __init__(self):super().__init__();self.refs=[]
            def handle_starttag(self,tag,attrs):
                for key,value in attrs:
                    if key in ('src','href') and value and not value.startswith(('#','http')):self.refs.append(value.split('#')[0])
        parser=References();parser.feed((ROOT/'ml-model-comparison.html').read_text(encoding='utf-8'))
        for ref in parser.refs:
            self.assertTrue((ROOT/ref).is_file(),ref)
            self.assertFalse(ref.startswith(('docs/','tests/','scripts/')),ref)
        config=json.loads((ROOT/'vercel.json').read_text())
        self.assertEqual(set(config),{'$schema','functions'})
        self.assertEqual(set(config['functions']),{'api/ml-compare.py'})
        self.assertEqual(config['functions']['api/ml-compare.py']['maxDuration'],60)
        js=(ROOT/'ml-model-comparison.js').read_text(encoding='utf-8')
        self.assertNotIn('http://localhost',js);self.assertNotIn('127.0.0.1',js)
        self.assertTrue(issubclass(lab.handler,lab.BaseHTTPRequestHandler))
        for name in ['api/market-data.js','data-analyzer.js','health-prediction.js','customer-return.js','vietnam-estimate-core.js','auth.js','history.js','supabase-config.js']:
            self.assertEqual((ROOT/name).read_bytes().replace(b'\r\n',b'\n'),subprocess.check_output(['git','show','HEAD:'+name],cwd=ROOT).replace(b'\r\n',b'\n'),name)


def browser_checks():
    from test_vietnam_map_online import CDP
    browser=os.environ.get('HEALTH_TEST_BROWSER',r'C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe')
    if not Path(browser).exists():raise RuntimeError('Set HEALTH_TEST_BROWSER to Chromium/Edge.')
    output=Path(tempfile.gettempdir())/'ml-model-comparison-qa'
    output.mkdir(exist_ok=True)
    scripts=(ROOT/'tests/ml-model-comparison-browser.js').read_text(encoding='utf-8')
    class Local(lab.handler,SimpleHTTPRequestHandler):
        def do_GET(self):
            if self.path=='/api/ml-compare':return lab.handler.do_GET(self)
            return SimpleHTTPRequestHandler.do_GET(self)
        def do_POST(self):
            if self.path!='/api/ml-compare':return self.respond(404,{'error':'Not found'})
            return lab.handler.do_POST(self)
    import functools
    server=ThreadingHTTPServer(('127.0.0.1',0),functools.partial(Local,directory=str(ROOT)))
    threading.Thread(target=server.serve_forever,daemon=True).start()
    process=None
    reports=[]
    try:
        with tempfile.TemporaryDirectory(prefix='ml-comparison-browser-',ignore_cleanup_errors=True) as profile:
            process=subprocess.Popen([browser,'--headless=new','--disable-gpu','--no-first-run','--disable-extensions','--disable-background-networking','--remote-debugging-port=0','--remote-allow-origins=*','--user-data-dir='+profile,'about:blank'],stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL,creationflags=subprocess.CREATE_NO_WINDOW if os.name=='nt' else 0)
            portfile=Path(profile)/'DevToolsActivePort'
            for _ in range(150):
                if portfile.exists():break
                time.sleep(.1)
            port=int(portfile.read_text().splitlines()[0])
            target=json.load(urllib.request.urlopen(urllib.request.Request(f'http://127.0.0.1:{port}/json/new?about:blank',method='PUT')))
            cdp=CDP(target['webSocketDebuggerUrl']);cdp.s.settimeout(180);cdp.call('Page.enable')
            def js(code):
                response=cdp.call('Runtime.evaluate',{'expression':code,'returnByValue':True,'awaitPromise':True})
                if 'exceptionDetails' in response:raise RuntimeError(str(response['exceptionDetails']))
                return response.get('result',{}).get('value')
            for width in [320,390,768,1024,1440]:
                cdp.call('Emulation.setDeviceMetricsOverride',{'width':width,'height':1000,'deviceScaleFactor':1,'mobile':width<=390})
                cdp.call('Page.navigate',{'url':f'http://127.0.0.1:{server.server_port}/ml-model-comparison.html'})
                for _ in range(150):
                    if js("document.readyState==='complete'&&typeof MLComparison!=='undefined'"):break
                    time.sleep(.05)
                js('window.__mlFixtures='+json.dumps({'classification':fixture(),'regression':fixture('regression'), 'excel':excel_fixture(multiple=True),'weak':weak_fixture()}))
                report=js(scripts);reports.append(report);print(json.dumps(report),flush=True)
                js('scrollTo(0,0)')
                screenshot=cdp.call('Page.captureScreenshot',{'format':'png'})
                (output/f'page-{width}.png').write_bytes(base64.b64decode(screenshot['data']))
                y=js("document.getElementById('results').getBoundingClientRect().top+scrollY")
                screenshot=cdp.call('Page.captureScreenshot',{'format':'png','captureBeyondViewport':True,'clip':{'x':0,'y':max(0,y),'width':width,'height':1800,'scale':1}})
                (output/f'results-{width}.png').write_bytes(base64.b64decode(screenshot['data']))
            previous=js("['modelA','modelB','detailModel'].map(id=>document.getElementById(id).value)")
            cdp.call('Emulation.setDeviceMetricsOverride',{'width':320,'height':1000,'deviceScaleFactor':1,'mobile':True})
            time.sleep(.3)
            assert js("document.documentElement.scrollWidth<=innerWidth && document.querySelector('#comparisonChart svg').viewBox.baseVal.width<=280")
            assert js("['modelA','modelB','detailModel'].map(id=>document.getElementById(id).value)")==previous
            assert js('window.__mlTestErrors.length')==0
            print('PASS live resize: charts remain readable; model selections preserved.',flush=True)
            cdp.call('Browser.close')
            process.wait(timeout=15)
    finally:
        if process and process.poll() is None:
            if os.name=='nt':subprocess.run(['taskkill','/PID',str(process.pid),'/T','/F'],stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL,creationflags=subprocess.CREATE_NO_WINDOW)
            else:process.kill()
        server.shutdown();server.server_close()
    (output/'report.json').write_text(json.dumps(reports,indent=2),encoding='utf-8')
    print('PASS browser: five viewports. Screenshots:',output)


if __name__=='__main__':
    suite=unittest.TestSuite([unittest.defaultTestLoader.loadTestsFromTestCase(cls) for cls in (Backend,HTTP)])
    result=unittest.TextTestRunner(verbosity=2).run(suite)
    if not result.wasSuccessful():sys.exit(1)
    if '--backend-only' not in sys.argv:browser_checks()
